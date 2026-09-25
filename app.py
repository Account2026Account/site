import os, random, sqlite3
import csv, io
from datetime import date, datetime, timedelta
from functools import wraps
from flask import Flask, Response, g, render_template, request, redirect, url_for, session, jsonify, flash, abort
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename

_ROOT = os.path.dirname(os.path.abspath(__file__))
_BASE = os.path.dirname(_ROOT)  # campuse_site/
DB = os.path.join(_ROOT, "campus.db")
UPLOADS = os.path.join(_BASE, "uploads")
os.makedirs(UPLOADS, exist_ok=True)
ALLOWED_EXT = {".pdf", ".doc", ".docx", ".xls", ".xlsx", ".ppt", ".pptx", ".txt", ".png", ".jpg", ".jpeg", ".gif", ".webp", ".zip", ".rar"}
app = Flask(__name__, template_folder=os.path.join(_BASE, "templates"), static_folder=os.path.join(_BASE, "static"))
app.secret_key = "123"  # замените перед публикацией

# Связи: преподаватель -> назначение (группа + предмет); занятия и отметки общие для группы и предмета
SCHEMA = """
CREATE TABLE groups(id INTEGER PRIMARY KEY, name TEXT UNIQUE NOT NULL);
CREATE TABLE subjects(id INTEGER PRIMARY KEY, name TEXT UNIQUE NOT NULL);
CREATE TABLE users(id INTEGER PRIMARY KEY, login TEXT UNIQUE NOT NULL, pw TEXT NOT NULL, name TEXT NOT NULL,
  role TEXT NOT NULL, group_id INTEGER REFERENCES groups(id));
CREATE TABLE assignments(id INTEGER PRIMARY KEY,
  teacher_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
  group_id INTEGER NOT NULL REFERENCES groups(id) ON DELETE CASCADE,
  subject_id INTEGER NOT NULL REFERENCES subjects(id) ON DELETE CASCADE, UNIQUE(teacher_id, group_id, subject_id));
CREATE TABLE lessons(id INTEGER PRIMARY KEY,
  group_id INTEGER NOT NULL REFERENCES groups(id) ON DELETE CASCADE,
  subject_id INTEGER NOT NULL REFERENCES subjects(id) ON DELETE CASCADE,
  day TEXT NOT NULL, topic TEXT DEFAULT '', teacher_id INTEGER REFERENCES users(id) ON DELETE SET NULL,
  UNIQUE(group_id, subject_id, day));
CREATE TABLE marks(student_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
  lesson_id INTEGER NOT NULL REFERENCES lessons(id) ON DELETE CASCADE,
  present INTEGER, grade INTEGER, PRIMARY KEY(student_id, lesson_id));
CREATE TABLE announcements(id INTEGER PRIMARY KEY, author_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
  group_id INTEGER REFERENCES groups(id) ON DELETE CASCADE, title TEXT NOT NULL, body TEXT NOT NULL, created TEXT NOT NULL);
"""

def init_db():
    if os.path.exists(DB):
        return False
    c = sqlite3.connect(DB); c.execute("PRAGMA foreign_keys=ON"); c.executescript(SCHEMA)
    def add(login, pw, name, role, gid=None):
        return c.execute("INSERT INTO users(login,pw,name,role,group_id) VALUES(?,?,?,?,?)",
                         (login, generate_password_hash(pw), name, role, gid)).lastrowid
    a = add("admin", "admin123", "Администратор", "admin")
    t1 = add("teacher1", "teacher123", "Анна Петрова", "teacher")
    t2 = add("teacher2", "teacher123", "Сергей Сидоров", "teacher")
    for n in ("ИС-21", "ПИ-22"): c.execute("INSERT INTO groups(name) VALUES(?)", (n,))
    for n in ("Математика", "Программирование", "Физика"): c.execute("INSERT INTO subjects(name) VALUES(?)", (n,))
    names = {1: ["Иван Смирнов", "Мария Кузнецова", "Дмитрий Орлов"], 2: ["Елена Волкова", "Артём Лебедев", "София Морозова"]}
    ids, k = {1: [], 2: []}, 0
    for gid, ns in names.items():
        for n in ns:
            k += 1; ids[gid].append(add(f"student{k}", "student123", n, "student", gid))
    for t, gid, sub in [(t1, 1, 1), (t2, 1, 1), (t1, 1, 2), (t2, 1, 3), (t1, 2, 1), (t2, 2, 2), (t2, 2, 3), (t1, 2, 3)]:
        c.execute("INSERT INTO assignments(teacher_id,group_id,subject_id) VALUES(?,?,?)", (t, gid, sub))
    now = datetime.now().strftime("%d.%m.%Y %H:%M")
    c.execute("INSERT INTO announcements(author_id,group_id,title,body,created) VALUES(?,?,?,?,?)", (a, None, "Добро пожаловать в Campus", "Оценки и посещаемость теперь доступны онлайн.", now))
    c.execute("INSERT INTO announcements(author_id,group_id,title,body,created) VALUES(?,?,?,?,?)", (t1, 1, "Контрольная по математике", "В пятницу принесите калькуляторы.", now))
    c.commit(); c.close()
    return True

def db():
    if "db" not in g:
        g.db = sqlite3.connect(DB); g.db.row_factory = sqlite3.Row
        g.db.execute("PRAGMA foreign_keys=ON")
    return g.db

@app.teardown_appcontext
def close_db(_):
    d = g.pop("db", None)
    if d: d.close()

def role(*rs):
    def deco(f):
        @wraps(f)
        def wrap(*a, **k):
            if session.get("role") not in rs:
                return redirect(url_for("login"))
            return f(*a, **k)
        return wrap
    return deco

def stats(cells):
    """cells: строки с полями present/grade (или None). Возвращает средний балл, пропуски, посещаемость."""
    gr = [c["grade"] for c in cells if c and c["grade"]]
    rec = [c for c in cells if c and c["present"] is not None]
    came = sum(1 for c in rec if c["present"])
    return dict(avg=round(sum(gr) / len(gr), 2) if gr else None, missed=len(rec) - came,
                pct=round(came * 100 / len(rec)) if rec else None)

@app.route("/")
def index():
    r = session.get("role")
    return redirect(url_for(r) if r else url_for("login"))

@app.route("/login", methods=["GET", "POST"])
def login():
    err = None
    if request.method == "POST":
        u = db().execute("SELECT * FROM users WHERE login=?", (request.form["login"].strip(),)).fetchone()
        if u and check_password_hash(u["pw"], request.form["password"]):
            session.clear(); session.update(uid=u["id"], role=u["role"], name=u["name"])
            return redirect(url_for("index"))
        err = "Неверный логин или пароль. Проверьте данные и попробуйте снова."
    return render_template("login.html", err=err)

@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))

WD = ["Пн", "Вт", "Ср", "Чт", "Пт", "Сб"]
WDF = ["Понедельник", "Вторник", "Среда", "Четверг", "Пятница", "Суббота"]

class Bad(Exception):
    pass

# Расписание: звонки, недельный шаблон (в т.ч. чётные/нечётные недели и разовые занятия), замены и отмены
NEW_SCHEMA = """
CREATE TABLE IF NOT EXISTS bells(id INTEGER PRIMARY KEY, num INTEGER UNIQUE NOT NULL, start TEXT NOT NULL, end TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS schedule(id INTEGER PRIMARY KEY,
  group_id INTEGER NOT NULL REFERENCES groups(id) ON DELETE CASCADE,
  subject_id INTEGER NOT NULL REFERENCES subjects(id) ON DELETE CASCADE,
  teacher_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
  bell_id INTEGER NOT NULL REFERENCES bells(id) ON DELETE CASCADE,
  weekday INTEGER, day TEXT, room TEXT DEFAULT '', parity INTEGER NOT NULL DEFAULT 0);
CREATE TABLE IF NOT EXISTS changes(id INTEGER PRIMARY KEY,
  schedule_id INTEGER NOT NULL REFERENCES schedule(id) ON DELETE CASCADE, day TEXT NOT NULL, kind TEXT NOT NULL,
  teacher_id INTEGER REFERENCES users(id) ON DELETE SET NULL, room TEXT DEFAULT '', note TEXT DEFAULT '', UNIQUE(schedule_id, day));
CREATE TABLE IF NOT EXISTS settings(key TEXT PRIMARY KEY, value TEXT);
CREATE TABLE IF NOT EXISTS duties(id INTEGER PRIMARY KEY, group_id INTEGER NOT NULL REFERENCES groups(id) ON DELETE CASCADE, day TEXT NOT NULL,
  student_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE, task TEXT NOT NULL DEFAULT 'Дежурный по группе', UNIQUE(group_id, day, student_id));
CREATE TABLE IF NOT EXISTS homework(id INTEGER PRIMARY KEY, group_id INTEGER NOT NULL REFERENCES groups(id) ON DELETE CASCADE,
  subject_id INTEGER NOT NULL REFERENCES subjects(id) ON DELETE CASCADE, author_id INTEGER REFERENCES users(id) ON DELETE SET NULL,
  title TEXT NOT NULL, body TEXT DEFAULT '', due TEXT, created TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS hw_done(hw_id INTEGER NOT NULL REFERENCES homework(id) ON DELETE CASCADE,
  student_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE, PRIMARY KEY(hw_id, student_id));
CREATE TABLE IF NOT EXISTS hw_attach(id INTEGER PRIMARY KEY,
  hw_id INTEGER NOT NULL REFERENCES homework(id) ON DELETE CASCADE,
  kind TEXT NOT NULL, name TEXT NOT NULL, url TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS hw_notes(hw_id INTEGER NOT NULL REFERENCES homework(id) ON DELETE CASCADE,
  student_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
  body TEXT DEFAULT '', updated TEXT NOT NULL, PRIMARY KEY(hw_id, student_id));
"""

def setting(c, key):
    r = c.execute("SELECT value FROM settings WHERE key=?", (key,)).fetchone()
    return r[0] if r else None

def minutes(t):
    h, m = map(int, t.split(":"))
    return h * 60 + m

def gap_label(m):
    return f"перерыв {m} мин" if m <= 40 else f"окно {m // 60} ч {m % 60} мин" if m % 60 else f"окно {m // 60} ч"

def overlap(a, b):
    """Пересекаются ли две записи расписания (по дню недели, чётности или конкретной дате)."""
    if a["day"] and b["day"]:
        return a["day"] == b["day"]
    if a["day"] or b["day"]:
        o, w = (a, b) if a["day"] else (b, a)
        dd = date.fromisoformat(o["day"])
        return w["weekday"] == dd.weekday() and w["parity"] in (0, 1 if dd.isocalendar()[1] % 2 else 2)
    return a["weekday"] == b["weekday"] and (a["parity"] == 0 or b["parity"] == 0 or a["parity"] == b["parity"])

def occurrences(c, start, end, gid=None, tid=None):
    """Реальные занятия за период: шаблон недели + разовые занятия + замены и отмены."""
    bells = {b["id"]: b for b in c.execute("SELECT * FROM bells")}
    rows = c.execute("""SELECT s.*, g.name gname, sub.name sname, t.name tname FROM schedule s
      JOIN groups g ON g.id=s.group_id JOIN subjects sub ON sub.id=s.subject_id JOIN users t ON t.id=s.teacher_id""").fetchall()
    ch = {(x["schedule_id"], x["day"]): x for x in c.execute(
        "SELECT c.*, u.name tname FROM changes c LEFT JOIN users u ON u.id=c.teacher_id")}
    out, d = [], start
    while d <= end:
        iso, par = d.isoformat(), 1 if d.isocalendar()[1] % 2 else 2
        for s in rows:
            if s["day"]:
                if s["day"] != iso: continue
            elif s["weekday"] != d.weekday() or s["parity"] not in (0, par):
                continue
            if gid and s["group_id"] != gid: continue
            x, b = ch.get((s["id"], iso)), bells[s["bell_id"]]
            rep_ = bool(x and x["kind"] == "replace")
            eff = x["teacher_id"] if rep_ and x["teacher_id"] else s["teacher_id"]
            if tid and tid not in (s["teacher_id"], eff): continue
            out.append(dict(id=s["id"], day=iso, weekday=d.weekday(), num=b["num"], start=b["start"], end=b["end"],
                gid=s["group_id"], sid=s["subject_id"], gname=s["gname"], sname=s["sname"], tid=eff,
                tname=x["tname"] if rep_ and x["tname"] else s["tname"], orig_tname=s["tname"],
                room=x["room"] if rep_ and x["room"] else s["room"], status=x["kind"] if x else "ok",
                note=x["note"] if x else "", chid=x["id"] if x else None))
        d += timedelta(days=1)
    out.sort(key=lambda o: (o["day"], o["start"]))
    return out

def sync_lessons(c, gid=None):
    """Создаёт занятия в журналах по расписанию (с начала семестра по сегодня), отменённые пропускает."""
    ts = setting(c, "term_start")
    start, end = date.fromisoformat(ts) if ts else date.today(), date.today()
    if start > end: return
    for o in occurrences(c, start, end, gid=gid):
        if o["status"] != "cancel":
            c.execute("""INSERT INTO lessons(group_id,subject_id,day,teacher_id) VALUES(?,?,?,?)
              ON CONFLICT(group_id,subject_id,day) DO UPDATE SET teacher_id=excluded.teacher_id""",
                      (o["gid"], o["sid"], o["day"], o["tid"]))
    c.commit()

def migrate():
    """Добавляет таблицы расписания. Работает и на старой базе, данные не теряются."""
    c = sqlite3.connect(DB); c.row_factory = sqlite3.Row; c.execute("PRAGMA foreign_keys=ON"); c.executescript(NEW_SCHEMA)
    if "head_id" not in [r[1] for r in c.execute("PRAGMA table_info(groups)")]:  # староста группы
        c.execute("ALTER TABLE groups ADD COLUMN head_id INTEGER REFERENCES users(id) ON DELETE SET NULL")
    if not setting(c, "v3"):
        if not c.execute("SELECT 1 FROM bells").fetchone():
            c.executemany("INSERT INTO bells(num,start,end) VALUES(?,?,?)", [(1, "08:30", "10:00"), (2, "10:10", "11:40"),
                (3, "12:10", "13:40"), (4, "13:50", "15:20"), (5, "15:30", "17:00"), (6, "17:10", "18:40")])
        c.execute("INSERT OR IGNORE INTO users(login,pw,name,role) VALUES('dispatcher',?,'Ольга Диспетчерова','dispatcher')",
                  (generate_password_hash("dispatcher123"),))
        c.execute("INSERT OR REPLACE INTO settings VALUES('term_start',?)", ((date.today() - timedelta(days=28)).isoformat(),))
        bells = [b["id"] for b in c.execute("SELECT id FROM bells ORDER BY num")]
        for a in c.execute("SELECT group_id, subject_id, MIN(teacher_id) t FROM assignments GROUP BY group_id, subject_id").fetchall():
            g_, s_ = a["group_id"], a["subject_id"]
            for wd in (s_ % 5, (s_ + 2) % 5):  # демо-расписание по существующим назначениям
                c.execute("INSERT INTO schedule(group_id,subject_id,teacher_id,bell_id,weekday,room) VALUES(?,?,?,?,?,?)",
                          (g_, s_, a["t"], bells[(s_ + g_ - 2) % len(bells)], wd, str(100 + s_ * 10 + g_)))
        c.execute("INSERT OR REPLACE INTO settings VALUES('v3','1')")
    sync_lessons(c)
    c.close()

def seed_marks():
    c = sqlite3.connect(DB); c.row_factory = sqlite3.Row; random.seed(1)
    studs, cnt = {}, {}
    for u in c.execute("SELECT id, group_id FROM users WHERE role='student'"):
        studs.setdefault(u["group_id"], []).append(u["id"])
    topics = ["Введение", "Основные понятия", "Практическое занятие", "Разбор задач", "Контрольная работа"]
    for l in c.execute("SELECT id, group_id, subject_id FROM lessons ORDER BY group_id, subject_id, day").fetchall():
        k = (l["group_id"], l["subject_id"]); n = cnt[k] = cnt.get(k, -1) + 1
        c.execute("UPDATE lessons SET topic=? WHERE id=?", (topics[n % 5], l["id"]))
        for s in studs.get(l["group_id"], []):
            p = int(random.random() < .85)
            c.execute("INSERT INTO marks VALUES(?,?,?,?)", (s, l["id"], p, random.choice([2, 3, 4, 4, 5, 5]) if p and random.random() < .6 else None))
    for g in c.execute("SELECT id FROM groups").fetchall():  # демо: староста, график дежурств
        first = c.execute("SELECT id FROM users WHERE role='student' AND group_id=? ORDER BY id LIMIT 1", (g["id"],)).fetchone()
        if first:
            c.execute("UPDATE groups SET head_id=? WHERE id=?", (first["id"], g["id"]))
            gen_duty(c, g["id"], date.today(), 10, 1, "Дежурный по группе")
    for a in c.execute("SELECT group_id, subject_id, MIN(teacher_id) t FROM assignments GROUP BY group_id, subject_id").fetchall():
        c.execute("INSERT INTO homework(group_id,subject_id,author_id,title,body,due,created) VALUES(?,?,?,?,?,?,?)",
                  (a["group_id"], a["subject_id"], a["t"], "Практическая работа №1", "Выполнить задания из методички.",
                   (date.today() + timedelta(days=5)).isoformat(), date.today().isoformat()))
    c.commit(); c.close()

def setup():
    fresh = init_db()
    migrate()
    if fresh: seed_marks()

def gen_duty(c, gid, start, days, per, task):
    """Автоматический график дежурств: по кругу по алфавиту, продолжая с того, кто дежурил последним."""
    studs = [r[0] for r in c.execute("SELECT id FROM users WHERE role='student' AND group_id=? ORDER BY name", (gid,))]
    if not studs: return
    last = c.execute("SELECT student_id FROM duties WHERE group_id=? AND day<? ORDER BY day DESC, id DESC LIMIT 1", (gid, start.isoformat())).fetchone()
    i = (studs.index(last[0]) + 1) % len(studs) if last and last[0] in studs else 0
    n, day = 0, start
    while n < days:
        if day.weekday() < 5:
            for _ in range(min(per, len(studs))):
                c.execute("INSERT OR IGNORE INTO duties(group_id,day,student_id,task) VALUES(?,?,?,?)", (gid, day.isoformat(), studs[i % len(studs)], task))
                i += 1
            n += 1
        day += timedelta(days=1)

def feed(c, gid=None, tid=None):
    t = date.today()
    week = occurrences(c, t, t + timedelta(days=7), gid=gid, tid=tid)
    return dict(today_items=[o for o in week if o["day"] == t.isoformat()], changes_items=[o for o in week if o["status"] != "ok"])

@app.context_processor
def inject():
    return dict(now=datetime.now(), WDF=WDF)

# ---------- Общее ----------
def announcements(where="", args=()):
    return db().execute("""SELECT a.*, u.name author, g.name gname FROM announcements a JOIN users u ON u.id=a.author_id
      LEFT JOIN groups g ON g.id=a.group_id """ + where + " ORDER BY a.id DESC LIMIT 8", args).fetchall()

@app.post("/announce")
def announce():
    r, d, f = session.get("role"), db(), request.form
    if r not in ("teacher", "admin"): abort(403)
    gid = f.get("group_id") or None
    if r == "teacher" and not (gid and d.execute("SELECT 1 FROM assignments WHERE teacher_id=? AND group_id=?", (session["uid"], gid)).fetchone()):
        abort(403)
    title, body = f["title"].strip(), f["body"].strip()
    if title and body:
        d.execute("INSERT INTO announcements(author_id,group_id,title,body,created) VALUES(?,?,?,?,?)",
                  (session["uid"], gid, title, body, datetime.now().strftime("%d.%m.%Y %H:%M")))
        d.commit(); flash("Объявление опубликовано")
    return redirect(url_for("index"))

@app.post("/announce/delete")
def announce_del():
    if session.get("role") not in ("teacher", "admin"): abort(403)
    d = db()
    d.execute("DELETE FROM announcements WHERE id=? AND (author_id=? OR ?='admin')", (request.form["id"], session["uid"], session["role"]))
    d.commit()
    return redirect(url_for("index"))

@app.route("/profile", methods=["GET", "POST"])
def profile():
    if "uid" not in session: return redirect(url_for("login"))
    if request.method == "POST":
        d, f = db(), request.form
        u = d.execute("SELECT pw FROM users WHERE id=?", (session["uid"],)).fetchone()
        if not check_password_hash(u["pw"], f["old"]): flash("Текущий пароль указан неверно")
        elif len(f["new"]) < 6: flash("Новый пароль должен быть не короче 6 символов")
        elif f["new"] != f["again"]: flash("Пароли не совпадают")
        else:
            d.execute("UPDATE users SET pw=? WHERE id=?", (generate_password_hash(f["new"]), session["uid"])); d.commit()
            flash("Пароль изменён")
        return redirect(url_for("profile"))
    return render_template("profile.html")

# ---------- Преподаватель ----------
@app.route("/teacher")
@role("teacher")
def teacher():
    d, uid = db(), session["uid"]
    sync_lessons(d)
    rows = d.execute("""SELECT g.id gid, s.id sid, g.name gname, s.name sname,
      (SELECT COUNT(*) FROM users WHERE group_id=g.id AND role='student') n,
      (SELECT COUNT(*) FROM lessons WHERE group_id=g.id AND subject_id=s.id) l,
      (SELECT group_concat(t.name, ', ') FROM assignments a2 JOIN users t ON t.id=a2.teacher_id
        WHERE a2.group_id=g.id AND a2.subject_id=s.id AND a2.teacher_id<>?) co
      FROM assignments a JOIN groups g ON g.id=a.group_id JOIN subjects s ON s.id=a.subject_id
      WHERE a.teacher_id=? ORDER BY g.name, s.name""", (uid, uid)).fetchall()
    groups = d.execute("SELECT DISTINCT g.id, g.name FROM assignments a JOIN groups g ON g.id=a.group_id WHERE a.teacher_id=? ORDER BY g.name", (uid,)).fetchall()
    return render_template("teacher.html", rows=rows, groups=groups, anns=announcements("WHERE a.author_id=?", (uid,)), can_del=True,
                           **feed(d, tid=uid))

def course(gid, sid):
    c = db().execute("""SELECT g.id gid, s.id sid, g.name gname, s.name sname FROM assignments a
      JOIN groups g ON g.id=a.group_id JOIN subjects s ON s.id=a.subject_id
      WHERE a.group_id=? AND a.subject_id=? AND a.teacher_id=?""", (gid, sid, session["uid"])).fetchone()
    if not c: abort(404)
    return c

def load_journal(gid, sid):
    d = db(); sync_lessons(d, gid)
    lessons = d.execute("""SELECT l.id, l.day, l.topic, t.name tname FROM lessons l LEFT JOIN users t ON t.id=l.teacher_id
      WHERE l.group_id=? AND l.subject_id=? ORDER BY l.day""", (gid, sid)).fetchall()
    studs = d.execute("SELECT id, name FROM users WHERE role='student' AND group_id=? ORDER BY name", (gid,)).fetchall()
    m = {(r["student_id"], r["lesson_id"]): r for r in d.execute(
        "SELECT * FROM marks WHERE lesson_id IN (SELECT id FROM lessons WHERE group_id=? AND subject_id=?)", (gid, sid))}
    rows = []
    for s in studs:
        cells = [m.get((s["id"], l["id"])) for l in lessons]
        rows.append(dict(id=s["id"], name=s["name"], cells=cells, **stats(cells)))
    return lessons, rows

@app.route("/journal/<int:gid>/<int:sid>")
@role("teacher")
def journal(gid, sid):
    c = course(gid, sid)
    lessons, rows = load_journal(gid, sid)
    team = ", ".join(r["name"] for r in db().execute("""SELECT t.name FROM assignments a JOIN users t ON t.id=a.teacher_id
      WHERE a.group_id=? AND a.subject_id=? ORDER BY t.name""", (gid, sid)))
    return render_template("journal.html", a=c, lessons=lessons, rows=rows, team=team, today=date.today().isoformat())

@app.route("/journal/<int:gid>/<int:sid>/export")
@role("teacher")
def export(gid, sid):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter as L
    c = course(gid, sid)
    lessons, rows = load_journal(gid, sid)
    team = ", ".join(r["name"] for r in db().execute("""SELECT t.name FROM assignments a JOIN users t ON t.id=a.teacher_id
      WHERE a.group_id=? AND a.subject_id=? ORDER BY t.name""", (gid, sid)))
    wb = Workbook(); ws = wb.active; ws.title = "Журнал"; ws.sheet_view.showGridLines = False
    fill = lambda h: PatternFill("solid", fgColor=h)
    thin = Side(style="thin", color="D9DEEC"); box = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws["A1"] = f'{c["gname"]}, {c["sname"]}'; ws["A1"].font = Font(size=16, bold=True, color="14213D")
    ws["A2"] = f"Преподаватели: {team}. Выгружено {datetime.now():%d.%m.%Y %H:%M}"; ws["A2"].font = Font(color="66708C")
    heads = ["Студент", "Средний балл", "Пропуски", "Посещаемость"] + [f'{l["day"][8:10]}.{l["day"][5:7]}' for l in lessons]
    for i, h in enumerate(heads, 1):
        x = ws.cell(4, i, h); x.font = Font(bold=True, color="FFFFFF"); x.fill = fill("14213D"); x.alignment = center; x.border = box
    ws.cell(5, 1, "Тема занятия")
    for i in range(1, 5 + len(lessons)):
        ws.cell(5, i).border = box; ws.cell(5, i).alignment = center; ws.cell(5, i).font = Font(italic=True, size=9, color="66708C")
    for i, l in enumerate(lessons, 5):
        ws.cell(5, i, l["topic"] or "")
    gfill = {2: "FFC9C9", 3: "FFD84D", 4: "CFE0FF", 5: "B8F0D3"}
    for ri, r in enumerate(rows, 6):
        risk = (r["avg"] and r["avg"] < 3) or (r["pct"] is not None and r["pct"] < 70)
        for i, v in enumerate([r["name"], r["avg"], r["missed"], r["pct"] / 100 if r["pct"] is not None else None], 1):
            x = ws.cell(ri, i, v); x.border = box
            x.alignment = Alignment(horizontal="left" if i == 1 else "center", vertical="center")
            if risk: x.fill = fill("FFF5F5")
        ws.cell(ri, 2).number_format = "0.00"; ws.cell(ri, 4).number_format = "0%"
        for j, cl in enumerate(r["cells"], 5):
            x = ws.cell(ri, j); x.border = box; x.alignment = center
            if cl and cl["grade"]:
                x.value = cl["grade"]; x.fill = fill(gfill[cl["grade"]]); x.font = Font(bold=True)
            elif cl and cl["present"] == 0:
                x.value = "н"; x.fill = fill("FFE4E4"); x.font = Font(bold=True, color="E5484D")
            elif cl and cl["present"] == 1:
                x.value = "•"; x.font = Font(color="1FA971")
    foot = 6 + len(rows)
    ws.cell(foot, 1, "Присутствовало")
    for k in range(len(lessons)):
        ws.cell(foot, 5 + k, sum(1 for r in rows if r["cells"][k] and r["cells"][k]["present"] == 1))
    for i in range(1, 5 + len(lessons)):
        x = ws.cell(foot, i); x.font = Font(bold=True); x.fill = fill("EEF1F8"); x.border = box; x.alignment = center
    ws.column_dimensions["A"].width = 30
    for i in range(2, 5): ws.column_dimensions[L(i)].width = 15
    for i in range(5, 5 + len(lessons)): ws.column_dimensions[L(i)].width = 12
    ws.row_dimensions[4].height = 28; ws.row_dimensions[5].height = 40
    ws.freeze_panes = "E6"; ws.page_setup.orientation = "landscape"
    buf = io.BytesIO(); wb.save(buf)
    return Response(buf.getvalue(), mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename=journal_{gid}_{sid}.xlsx"})

@app.post("/api/mark")
@role("teacher")
def mark():
    j, d = request.get_json(), db()
    l = d.execute("""SELECT l.id, l.group_id, l.subject_id FROM lessons l JOIN assignments a
      ON a.group_id=l.group_id AND a.subject_id=l.subject_id WHERE l.id=? AND a.teacher_id=?""", (j["lesson"], session["uid"])).fetchone()
    if not l: abort(403)
    if not d.execute("SELECT 1 FROM users WHERE id=? AND role='student' AND group_id=?", (j["student"], l["group_id"])).fetchone():
        abort(403)
    sid, lid = j["student"], l["id"]
    d.execute("INSERT OR IGNORE INTO marks(student_id,lesson_id) VALUES(?,?)", (sid, lid))
    if "present" in j:
        p = j["present"]
        d.execute("UPDATE marks SET present=? WHERE student_id=? AND lesson_id=?", (None if p is None else int(p), sid, lid))
        if p == 0 or p is None:  # отсутствовал или сброс: оценки быть не может
            d.execute("UPDATE marks SET grade=NULL WHERE student_id=? AND lesson_id=?", (sid, lid))
    if "grade" in j:
        if j["grade"] not in (None, 2, 3, 4, 5): abort(400)
        d.execute("UPDATE marks SET grade=? WHERE student_id=? AND lesson_id=?", (j["grade"], sid, lid))
        if j["grade"]:  # оценка означает, что студент был на занятии
            d.execute("UPDATE marks SET present=1 WHERE student_id=? AND lesson_id=?", (sid, lid))
    d.execute("DELETE FROM marks WHERE student_id=? AND lesson_id=? AND present IS NULL AND grade IS NULL", (sid, lid))
    d.commit()
    allm = d.execute("SELECT * FROM marks WHERE student_id=? AND lesson_id IN (SELECT id FROM lessons WHERE group_id=? AND subject_id=?)",
                     (sid, l["group_id"], l["subject_id"])).fetchall()
    cur = d.execute("SELECT present, grade FROM marks WHERE student_id=? AND lesson_id=?", (sid, lid)).fetchone()
    return jsonify(ok=True, present=cur["present"] if cur else None, grade=cur["grade"] if cur else None, **stats(allm))

@app.post("/api/lesson")
@role("teacher")
def lesson():
    j = request.get_json(); c = course(j["gid"], j["sid"]); d = db(); act = j["action"]
    if act == "add":
        try:
            d.execute("INSERT INTO lessons(group_id,subject_id,day,topic,teacher_id) VALUES(?,?,?,?,?)",
                      (c["gid"], c["sid"], date.fromisoformat(j["day"]).isoformat(), j.get("topic", "").strip(), session["uid"]))
        except ValueError:
            return jsonify(ok=False, error="Укажите корректную дату"), 400
        except sqlite3.IntegrityError:
            return jsonify(ok=False, error="Занятие на эту дату уже есть"), 409
    else:
        l = d.execute("SELECT id FROM lessons WHERE id=? AND group_id=? AND subject_id=?", (j["lesson"], c["gid"], c["sid"])).fetchone()
        if not l: abort(404)
        if act == "del":
            d.execute("DELETE FROM lessons WHERE id=?", (l["id"],))
        elif act == "topic":
            d.execute("UPDATE lessons SET topic=? WHERE id=?", (j["topic"].strip(), l["id"]))
        elif act == "all":  # всех без отметки посещения делаем присутствующими
            d.execute("INSERT OR IGNORE INTO marks(student_id,lesson_id) SELECT id,? FROM users WHERE role='student' AND group_id=?", (l["id"], c["gid"]))
            d.execute("UPDATE marks SET present=1 WHERE lesson_id=? AND present IS NULL", (l["id"],))
    d.commit()
    return jsonify(ok=True)

# ---------- Студент ----------
@app.route("/student")
@role("student")
def student():
    d, uid = db(), session["uid"]
    me = d.execute("SELECT u.group_id, g.name gname FROM users u LEFT JOIN groups g ON g.id=u.group_id WHERE u.id=?", (uid,)).fetchone()
    sync_lessons(d, me["group_id"])
    subs = []
    for a in d.execute("""SELECT s.id, s.name sname,
        (SELECT group_concat(t.name, ', ') FROM assignments a2 JOIN users t ON t.id=a2.teacher_id
          WHERE a2.group_id=? AND a2.subject_id=s.id) tname
        FROM subjects s WHERE s.id IN (SELECT subject_id FROM assignments WHERE group_id=?) ORDER BY s.name""",
                       (me["group_id"], me["group_id"])).fetchall():
        cells = d.execute("""SELECT l.day, l.topic, m.present, m.grade FROM lessons l
                             LEFT JOIN marks m ON m.lesson_id=l.id AND m.student_id=?
                             WHERE l.group_id=? AND l.subject_id=? ORDER BY l.day""", (uid, me["group_id"], a["id"])).fetchall()
        subs.append(dict(dict(a), cells=cells, **stats(cells)))
    allc = [c for s in subs for c in s["cells"]]
    tot = stats(allc)
    today = date.today().isoformat()
    my_duty = d.execute("SELECT day, task FROM duties WHERE student_id=? AND day>=? ORDER BY day LIMIT 3", (uid, today)).fetchall()
    hw = d.execute("""SELECT h.title, h.due, s.name sname FROM homework h JOIN subjects s ON s.id=h.subject_id WHERE h.group_id=?
      AND NOT EXISTS(SELECT 1 FROM hw_done x WHERE x.hw_id=h.id AND x.student_id=?) ORDER BY h.due LIMIT 3""", (me["group_id"], uid)).fetchall()
    return render_template("student.html", my_duty=my_duty, hw=hw, subs=subs, grp=me["gname"], avg=tot["avg"] or 0, pct=tot["pct"] or 0,
                           cnt=sum(1 for c in allc if c["grade"]), groups=None,
                           **feed(d, gid=me["group_id"] or -1),
                           anns=announcements("WHERE a.group_id IS NULL OR a.group_id=?", (me["group_id"],)))

# ---------- Администратор ----------
@app.route("/admin")
@role("admin")
def admin():
    d = db(); q = lambda x: d.execute(x).fetchall(); one = lambda x: d.execute(x).fetchone()[0]
    users = q("""SELECT u.*, g.name gname FROM users u LEFT JOIN groups g ON g.id=u.group_id
      ORDER BY CASE u.role WHEN 'admin' THEN 0 WHEN 'dispatcher' THEN 1 WHEN 'teacher' THEN 2 ELSE 3 END, g.name, u.name""")
    sbg = {}
    for u in users:
        if u["role"] == "student" and u["group_id"]: sbg.setdefault(u["group_id"], []).append(u)
    return render_template("admin.html", users=users, sbg=sbg, groups=q("SELECT * FROM groups ORDER BY name"),
        subjects=q("SELECT * FROM subjects ORDER BY name"), teachers=q("SELECT * FROM users WHERE role='teacher' ORDER BY name"),
        assigns=q("""SELECT a.id, g.name gname, s.name sname, t.name tname FROM assignments a JOIN groups g ON g.id=a.group_id
                     JOIN subjects s ON s.id=a.subject_id JOIN users t ON t.id=a.teacher_id ORDER BY g.name, s.name, t.name"""),
        counts=dict(groups=one("SELECT COUNT(*) FROM groups"), students=one("SELECT COUNT(*) FROM users WHERE role='student'"),
                    teachers=one("SELECT COUNT(*) FROM users WHERE role='teacher'"), lessons=one("SELECT COUNT(*) FROM lessons")),
        anns=announcements(), can_del=True, all_ok=True)

@app.post("/admin/<act>")
@role("admin")
def admin_act(act):
    d, f = db(), request.form
    try:
        if act == "group":
            d.execute("INSERT INTO groups(name) VALUES(?)", (f["name"].strip(),))
        elif act == "subject":
            d.execute("INSERT INTO subjects(name) VALUES(?)", (f["name"].strip(),))
        elif act == "user":
            r, gid = f["role"], f.get("group_id") or None
            if r not in ("teacher", "student", "dispatcher"): raise ValueError("Выберите роль")
            if r == "student" and not gid: raise ValueError("Для студента выберите группу")
            d.execute("INSERT INTO users(login,pw,name,role,group_id) VALUES(?,?,?,?,?)",
                      (f["login"].strip(), generate_password_hash(f["password"]), f["name"].strip(), r, gid if r == "student" else None))
        elif act == "assign":
            d.execute("INSERT INTO assignments(teacher_id,group_id,subject_id) VALUES(?,?,?)", (f["teacher_id"], f["group_id"], f["subject_id"]))
        elif act == "head":
            sid = f.get("student_id") or None
            if sid and not d.execute("SELECT 1 FROM users WHERE id=? AND group_id=? AND role='student'", (sid, f["group_id"])).fetchone():
                raise ValueError("Староста должен быть студентом этой группы")
            d.execute("UPDATE groups SET head_id=? WHERE id=?", (sid, f["group_id"]))
        elif act == "move":
            d.execute("UPDATE users SET group_id=? WHERE id=? AND role='student'", (f.get("group_id") or None, f["id"]))
        elif act == "pw":
            if len(f["password"]) < 6: raise ValueError("Пароль должен быть не короче 6 символов")
            d.execute("UPDATE users SET pw=? WHERE id=?", (generate_password_hash(f["password"]), f["id"]))
        elif act == "del":
            t = {"group": "groups", "subject": "subjects", "user": "users", "assign": "assignments"}[f["kind"]]
            if t == "users" and int(f["id"]) == session["uid"]: raise ValueError("Нельзя удалить самого себя")
            if t == "groups":  # студенты остаются в системе, но без группы
                d.execute("UPDATE users SET group_id=NULL WHERE group_id=?", (f["id"],))
            d.execute(f"DELETE FROM {t} WHERE id=?", (f["id"],))
        d.commit(); flash("Готово")
    except sqlite3.IntegrityError:
        d.rollback(); flash("Не получилось: такая запись уже есть, либо на неё ссылаются (например, в группе остались студенты).")
    except (ValueError, KeyError) as e:
        d.rollback(); flash(str(e))
    return redirect(url_for("admin"))

# ---------- Расписание ----------
def bell_rows(bells):
    out = []
    for i, b in enumerate(bells):
        r = dict(b)
        if i:
            g = minutes(b["start"]) - minutes(bells[i - 1]["end"])
            r["gap_label"] = gap_label(g) if g > 0 else None
        out.append(r)
    return out

def scope(d):
    r = session["role"]
    if r == "student":
        return d.execute("SELECT group_id FROM users WHERE id=?", (session["uid"],)).fetchone()[0] or -1, None, None
    if r == "teacher":
        # преподаватель видит только свои занятия; может отфильтровать по группе
        my_groups = d.execute("""SELECT DISTINCT g.id, g.name FROM assignments a
            JOIN groups g ON g.id=a.group_id WHERE a.teacher_id=? ORDER BY g.name""", (session["uid"],)).fetchall()
        gid = request.args.get("group", type=int)
        if gid and not any(g["id"] == gid for g in my_groups):
            gid = None
        return gid, session["uid"], my_groups
    return request.args.get("group", type=int), None, d.execute("SELECT * FROM groups ORDER BY name").fetchall()

@app.route("/schedule")
def schedule():
    if "uid" not in session: return redirect(url_for("login"))
    d, off = db(), request.args.get("week", 0, type=int)
    gid, tid, groups = scope(d)
    mon = date.today() - timedelta(days=date.today().weekday()) + timedelta(weeks=off)
    occ = occurrences(d, mon, mon + timedelta(days=5), gid=gid, tid=tid)
    days = []
    for i in range(6):
        day = mon + timedelta(days=i); lst = [o for o in occ if o["day"] == day.isoformat()]
        if gid or tid:
            for x, y in zip(lst, lst[1:]):
                g = minutes(y["start"]) - minutes(x["end"])
                if g > 0: y["gap_label"] = gap_label(g)
        days.append(dict(date=day, lst=lst))
    is_teacher = session.get("role") == "teacher"
    return render_template("schedule.html", days=days, off=off, gid=gid if gid and gid > 0 else None, groups=groups, wd=WDF,
                           today=date.today(), bells=bell_rows(d.execute("SELECT * FROM bells ORDER BY num").fetchall()),
                           is_teacher=is_teacher, show_group_filter=bool(groups))

@app.route("/schedule.xlsx")
def schedule_xlsx():
    """Экспорт расписания на неделю в Excel (сетка + список) с оформлением."""
    if "uid" not in session: return redirect(url_for("login"))
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter as L
    d, off = db(), request.args.get("week", 0, type=int)
    gid, tid, groups = scope(d)
    mon = date.today() - timedelta(days=date.today().weekday()) + timedelta(weeks=off)
    occ = occurrences(d, mon, mon + timedelta(days=5), gid=gid, tid=tid)
    bells = d.execute("SELECT * FROM bells ORDER BY num").fetchall()

    wb = Workbook()
    ws = wb.active
    ws.title = "Расписание"
    ws.sheet_view.showGridLines = False
    fill = lambda h: PatternFill("solid", fgColor=h)
    thin = Side(style="thin", color="D9DEEC")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="top", wrap_text=True)

    title = "Расписание"
    if gid and groups:
        gname = next((g["name"] for g in groups if g["id"] == gid), None)
        if gname:
            title += f" · {gname}"
    elif session.get("role") == "teacher":
        title += " · мои занятия"
    title += f"  {mon.strftime('%d.%m.%Y')} – {(mon + timedelta(days=5)).strftime('%d.%m.%Y')}"

    ws["A1"] = title
    ws["A1"].font = Font(size=16, bold=True, color="14213D")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=7)
    ws["A2"] = f"Выгружено {datetime.now():%d.%m.%Y %H:%M} · Campus"
    ws["A2"].font = Font(color="66708C", size=10)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=7)

    # Заголовки
    ws.cell(4, 1, "Пара / время").font = Font(bold=True, color="FFFFFF", size=11)
    ws.cell(4, 1).fill = fill("14213D")
    ws.cell(4, 1).alignment = center
    ws.cell(4, 1).border = box
    for i in range(6):
        cell = ws.cell(4, 2 + i, f"{WDF[i]}\n{(mon + timedelta(days=i)).strftime('%d.%m')}")
        cell.font = Font(bold=True, color="FFFFFF", size=11)
        cell.fill = fill("14213D")
        cell.alignment = center
        cell.border = box
    ws.row_dimensions[4].height = 36

    by_day_bell = {}
    for o in occ:
        by_day_bell.setdefault((o["day"], o["num"]), []).append(o)

    for ri, b in enumerate(bells, 5):
        time_cell = ws.cell(ri, 1, f'{b["num"]} пара\n{b["start"]}–{b["end"]}')
        time_cell.font = Font(bold=True, size=10)
        time_cell.alignment = center
        time_cell.border = box
        time_cell.fill = fill("EEF1F8")
        max_lines = 1
        for di in range(6):
            day = (mon + timedelta(days=di)).isoformat()
            items = by_day_bell.get((day, b["num"]), [])
            cell = ws.cell(ri, 2 + di)
            cell.border = box
            cell.alignment = left
            if not items:
                cell.value = ""
                continue
            blocks = []
            for o in items:
                st = ""
                if o["status"] == "cancel":
                    st = " [отменено]"
                elif o["status"] == "replace":
                    st = " [замена]"
                block = f'{o["sname"]}{st}\n{o["gname"]}\n{o["tname"]}\nауд. {o["room"] or "—"}'
                if o.get("note"):
                    block += f'\n{o["note"]}'
                blocks.append(block)
            text = "\n———\n".join(blocks)
            cell.value = text
            max_lines = max(max_lines, text.count("\n") + 1)
            if any(o["status"] == "cancel" for o in items):
                cell.fill = fill("FFE4E4")
            elif any(o["status"] == "replace" for o in items):
                cell.fill = fill("FFFBE8")
            else:
                cell.fill = fill("F5F7FE")
        ws.row_dimensions[ri].height = max(50, 14 * max_lines + 8)

    ws.column_dimensions["A"].width = 14
    for i in range(2, 8):
        ws.column_dimensions[L(i)].width = 22
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.freeze_panes = "B5"

    # Лист «Список»
    ws2 = wb.create_sheet("Список")
    ws2.sheet_view.showGridLines = False
    ws2["A1"] = title
    ws2["A1"].font = Font(size=14, bold=True, color="14213D")
    ws2.merge_cells("A1:J1")
    cols = ["Дата", "День", "Пара", "Время", "Предмет", "Группа", "Преподаватель", "Аудитория", "Статус", "Примечание"]
    for i, h in enumerate(cols, 1):
        x = ws2.cell(3, i, h)
        x.font = Font(bold=True, color="FFFFFF")
        x.fill = fill("14213D")
        x.alignment = center
        x.border = box
    for ri, o in enumerate(sorted(occ, key=lambda x: (x["day"], x["start"])), 4):
        st = {"ok": "ок", "cancel": "отменено", "replace": "замена"}.get(o["status"], o["status"])
        vals = [
            f'{o["day"][8:10]}.{o["day"][5:7]}.{o["day"][:4]}',
            WDF[o["weekday"]],
            o["num"],
            f'{o["start"]}–{o["end"]}',
            o["sname"],
            o["gname"],
            o["tname"],
            o["room"] or "—",
            st,
            o.get("note") or "",
        ]
        for ci, v in enumerate(vals, 1):
            x = ws2.cell(ri, ci, v)
            x.border = box
            x.alignment = center if ci <= 4 else left
            if o["status"] == "cancel":
                x.fill = fill("FFE4E4")
            elif o["status"] == "replace":
                x.fill = fill("FFFBE8")
    for i, w in enumerate([12, 10, 8, 12, 22, 12, 20, 12, 12, 24], 1):
        ws2.column_dimensions[L(i)].width = w
    if occ:
        ws2.auto_filter.ref = f"A3:J{3 + len(occ)}"
    ws2.freeze_panes = "A4"

    buf = io.BytesIO()
    wb.save(buf)
    fname = f"schedule_{mon.strftime('%Y%m%d')}.xlsx"
    return Response(
        buf.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={fname}"},
    )

@app.route("/dispatcher")
@role("dispatcher", "admin")
def dispatcher():
    d = db()
    groups = d.execute("SELECT * FROM groups ORDER BY name").fetchall()
    gid = request.args.get("group", type=int) or (groups[0]["id"] if groups else None)
    view = request.args.get("view") or "week"  # kept for compatibility, UI now shows both
    off = request.args.get("week", 0, type=int)
    mon = date.today() - timedelta(days=date.today().weekday()) + timedelta(weeks=off)
    grid, once, occ, wgrid = {}, [], [], {}
    if gid:
        for e in d.execute("""SELECT s.*, sub.name sname, t.name tname FROM schedule s JOIN subjects sub ON sub.id=s.subject_id
                JOIN users t ON t.id=s.teacher_id WHERE s.group_id=? ORDER BY s.parity""", (gid,)):
            if e["day"]: once.append(e)
            else: grid.setdefault((e["weekday"], e["bell_id"]), []).append(e)
        once.sort(key=lambda e: e["day"] or "")
        occ = occurrences(d, mon, mon + timedelta(days=5), gid=gid)
        for o in occ: wgrid.setdefault((o["weekday"], o["num"]), []).append(o)
    return render_template("dispatcher.html", groups=groups, gid=gid, view=view, off=off, grid=grid, once=once, occ=occ, wgrid=wgrid,
        dates=[mon + timedelta(days=i) for i in range(6)], wd=WD, term=setting(d, "term_start"),
        bells=bell_rows(d.execute("SELECT * FROM bells ORDER BY num").fetchall()),
        subjects=d.execute("SELECT * FROM subjects ORDER BY name").fetchall(),
        teachers=d.execute("SELECT id, name FROM users WHERE role='teacher' ORDER BY name").fetchall())

@app.post("/dispatcher/<act>")
@role("dispatcher", "admin")
def dispatcher_act(act):
    d, f = db(), request.form
    gid = f.get("group") or None
    try:
        if act == "bell":
            d.execute("INSERT INTO bells(num,start,end) VALUES(?,?,?)", (int(f["num"]), f["start"], f["end"]))
        elif act == "bell_del":
            d.execute("DELETE FROM bells WHERE id=?", (f["id"],))
        elif act == "term":
            d.execute("INSERT OR REPLACE INTO settings VALUES('term_start',?)", (date.fromisoformat(f["value"]).isoformat(),))
        elif act == "entry":
            new = dict(group_id=int(f["group"]), subject_id=int(f["subject_id"]), teacher_id=int(f["teacher_id"]), bell_id=int(f["bell_id"]),
                       weekday=int(f["weekday"]), day=f.get("day") or None, parity=int(f["parity"]))
            if new["day"]:
                new["day"] = date.fromisoformat(new["day"]).isoformat(); new["weekday"] = None
            for e in d.execute("SELECT * FROM schedule WHERE bell_id=?", (new["bell_id"],)):
                if (e["group_id"] == new["group_id"] or e["teacher_id"] == new["teacher_id"]) and overlap(e, new):
                    raise Bad("В это время у группы или у преподавателя уже есть занятие")
            d.execute("INSERT INTO schedule(group_id,subject_id,teacher_id,bell_id,weekday,day,room,parity) VALUES(?,?,?,?,?,?,?,?)",
                      (new["group_id"], new["subject_id"], new["teacher_id"], new["bell_id"], new["weekday"], new["day"], f.get("room", "").strip(), new["parity"]))
            d.execute("INSERT OR IGNORE INTO assignments(teacher_id,group_id,subject_id) VALUES(?,?,?)", (new["teacher_id"], new["group_id"], new["subject_id"]))
        elif act == "entry_del":
            d.execute("DELETE FROM schedule WHERE id=?", (f["id"],))
        elif act == "change":
            s = d.execute("SELECT * FROM schedule WHERE id=?", (f["schedule_id"],)).fetchone()
            if not s: raise Bad("Занятие не найдено")
            kind, day, room = f["kind"], date.fromisoformat(f["day"]).isoformat(), f.get("room", "").strip()
            tid = int(f["teacher_id"]) if f.get("teacher_id") else None
            if kind == "replace" and not (tid or room): raise Bad("Для замены укажите преподавателя или аудиторию")
            rep_ = kind == "replace"
            d.execute("""INSERT INTO changes(schedule_id,day,kind,teacher_id,room,note) VALUES(?,?,?,?,?,?)
              ON CONFLICT(schedule_id,day) DO UPDATE SET kind=excluded.kind, teacher_id=excluded.teacher_id, room=excluded.room, note=excluded.note""",
                      (s["id"], day, kind, tid if rep_ else None, room if rep_ else "", f.get("note", "").strip()))
            if kind == "cancel":  # пустое занятие убираем из журнала, с отметками оставляем
                d.execute("""DELETE FROM lessons WHERE group_id=? AND subject_id=? AND day=?
                  AND NOT EXISTS(SELECT 1 FROM marks WHERE lesson_id=lessons.id)""", (s["group_id"], s["subject_id"], day))
            elif tid:  # заменяющий преподаватель получает доступ к журналу группы
                d.execute("INSERT OR IGNORE INTO assignments(teacher_id,group_id,subject_id) VALUES(?,?,?)", (tid, s["group_id"], s["subject_id"]))
        elif act == "change_del":
            d.execute("DELETE FROM changes WHERE id=?", (f["id"],))
        d.commit(); flash("Готово")
    except Bad as e:
        d.rollback(); flash(str(e))
    except sqlite3.IntegrityError:
        d.rollback(); flash("Такая запись уже существует")
    except (ValueError, KeyError):
        d.rollback(); flash("Проверьте введённые данные")
    return redirect(url_for("dispatcher", group=gid, view=f.get("view") or None, week=f.get("week") or None))

# ---------- Дежурства (староста) ----------
def head_group(uid):
    r = db().execute("SELECT g.id FROM groups g JOIN users u ON u.group_id=g.id WHERE g.head_id=? AND u.id=?", (uid, uid)).fetchone()
    return r[0] if r else None

@app.route("/duty")
@role("student")
def duty():
    d, uid = db(), session["uid"]
    me = d.execute("""SELECT u.group_id, g.name gname, g.head_id, h.name hname FROM users u LEFT JOIN groups g ON g.id=u.group_id
      LEFT JOIN users h ON h.id=g.head_id WHERE u.id=?""", (uid,)).fetchone()
    gid, byday = me["group_id"], {}
    rows = d.execute("""SELECT d.*, u.name sname FROM duties d JOIN users u ON u.id=d.student_id WHERE d.group_id=? AND d.day>=?
      ORDER BY d.day, u.name""", (gid, (date.today() - timedelta(days=7)).isoformat())).fetchall() if gid else []
    for r in rows: byday.setdefault(r["day"], []).append(r)
    studs = d.execute("SELECT id, name FROM users WHERE role='student' AND group_id=? ORDER BY name", (gid,)).fetchall() if gid else []
    return render_template("duty.html", me=me, is_head=head_group(uid) is not None, studs=studs, wd=WD, today=date.today().isoformat(),
        duty_days=[dict(day=k, lst=v, wd=date.fromisoformat(k).weekday()) for k, v in byday.items()])

@app.post("/duty/<act>")
@role("student")
def duty_act(act):
    d, f = db(), request.form
    gid = head_group(session["uid"])
    if not gid: abort(403)
    task = f.get("task", "").strip() or "Дежурный по группе"
    try:
        if act == "add":
            if not d.execute("SELECT 1 FROM users WHERE id=? AND group_id=?", (f["student_id"], gid)).fetchone():
                raise Bad("Студент не из вашей группы")
            d.execute("INSERT INTO duties(group_id,day,student_id,task) VALUES(?,?,?,?)", (gid, date.fromisoformat(f["day"]).isoformat(), f["student_id"], task))
        elif act == "gen":
            gen_duty(d, gid, date.fromisoformat(f["start"]), max(1, min(int(f["days"]), 60)), max(1, min(int(f["per"]), 5)), task)
        elif act == "del":
            d.execute("DELETE FROM duties WHERE id=? AND group_id=?", (f["id"], gid))
        elif act == "clear":
            d.execute("DELETE FROM duties WHERE group_id=? AND day>=?", (gid, date.today().isoformat()))
        d.commit(); flash("Готово")
    except Bad as e:
        d.rollback(); flash(str(e))
    except sqlite3.IntegrityError:
        d.rollback(); flash("Этот студент уже назначен на этот день")
    except (ValueError, KeyError):
        d.rollback(); flash("Проверьте введённые данные")
    return redirect(url_for("duty"))

# ---------- Задания ----------

@app.route("/homework")
def homework():
    if "uid" not in session: return redirect(url_for("login"))
    d, uid, role = db(), session["uid"], session["role"]
    if role not in ("teacher", "student"): abort(403)
    courses = None
    if role == "teacher":
        courses = d.execute("""SELECT g.id gid, s.id sid, g.name gname, s.name sname FROM assignments a
            JOIN groups g ON g.id=a.group_id JOIN subjects s ON s.id=a.subject_id
            WHERE a.teacher_id=? ORDER BY g.name, s.name""", (uid,)).fetchall()
        rows = d.execute("""SELECT h.*, g.name gname, s.name sname,
          (SELECT COUNT(*) FROM hw_done x WHERE x.hw_id=h.id) done,
          (SELECT COUNT(*) FROM users u WHERE u.group_id=h.group_id AND u.role='student') total
          FROM homework h JOIN groups g ON g.id=h.group_id JOIN subjects s ON s.id=h.subject_id
          WHERE EXISTS(SELECT 1 FROM assignments a WHERE a.teacher_id=? AND a.group_id=h.group_id AND a.subject_id=h.subject_id)
          ORDER BY COALESCE(h.due, '9999'), h.id DESC""", (uid,)).fetchall()
        notes = {}
    else:
        gid = d.execute("SELECT group_id FROM users WHERE id=?", (uid,)).fetchone()[0]
        rows = d.execute("""SELECT h.*, s.name sname, u.name author,
          EXISTS(SELECT 1 FROM hw_done x WHERE x.hw_id=h.id AND x.student_id=?) done
          FROM homework h JOIN subjects s ON s.id=h.subject_id LEFT JOIN users u ON u.id=h.author_id
          WHERE h.group_id=? ORDER BY done, COALESCE(h.due, '9999'), h.id DESC""", (uid, gid)).fetchall()
        notes = {r["hw_id"]: r["body"] for r in d.execute(
            "SELECT hw_id, body FROM hw_notes WHERE student_id=?", (uid,))}
    # attachments for all listed homework
    ids = [r["id"] for r in rows]
    attach = {}
    if ids:
        q = ",".join("?" * len(ids))
        for a in d.execute(f"SELECT * FROM hw_attach WHERE hw_id IN ({q}) ORDER BY id", ids):
            attach.setdefault(a["hw_id"], []).append(a)
    # group by due date (or 'без срока')
    from collections import OrderedDict
    by_day = OrderedDict()
    today = date.today().isoformat()
    for r in rows:
        key = r["due"] or "без срока"
        by_day.setdefault(key, []).append(r)
    return render_template("homework.html", courses=courses, by_day=by_day, rows=rows,
                           today=today, attach=attach, notes=notes, is_teacher=(role == "teacher"))

@app.post("/homework/add")
@role("teacher")
def homework_add():
    f, d = request.form, db()
    try:
        gid, sid = map(int, f["course"].split(":"))
    except (ValueError, KeyError):
        abort(400)
    course(gid, sid)
    try:
        due = date.fromisoformat(f["due"]).isoformat() if f.get("due") else None
    except ValueError:
        due = None
    title = f["title"].strip()
    if not title:
        flash("Укажите название"); return redirect(url_for("homework"))
    cur = d.execute(
        "INSERT INTO homework(group_id,subject_id,author_id,title,body,due,created) VALUES(?,?,?,?,?,?,?)",
        (gid, sid, session["uid"], title, f.get("body", "").strip(), due, date.today().isoformat()))
    hw_id = cur.lastrowid
    # links (one per line or single field)
    for link in (f.get("links") or "").splitlines():
        link = link.strip()
        if not link:
            continue
        if not link.startswith(("http://", "https://")):
            link = "https://" + link
        name = f.get("link_name", "").strip() or link
        d.execute("INSERT INTO hw_attach(hw_id,kind,name,url) VALUES(?,?,?,?)", (hw_id, "link", name[:120], link[:500]))
    # files
    files = request.files.getlist("files")
    for fs in files:
        if not fs or not fs.filename:
            continue
        fn = secure_filename(fs.filename)
        ext = os.path.splitext(fn)[1].lower()
        if ext not in ALLOWED_EXT:
            continue
        if fs.content_length and fs.content_length > 15 * 1024 * 1024:
            continue
        stored = f"{hw_id}_{int(datetime.now().timestamp())}_{fn}"
        path = os.path.join(UPLOADS, stored)
        fs.save(path)
        d.execute("INSERT INTO hw_attach(hw_id,kind,name,url) VALUES(?,?,?,?)",
                  (hw_id, "file", fn, stored))
    d.commit()
    flash("Задание опубликовано")
    return redirect(url_for("homework"))

@app.post("/homework/delete")
@role("teacher")
def homework_del():
    d = db()
    hid = request.form["id"]
    row = d.execute("SELECT id FROM homework WHERE id=? AND author_id=?", (hid, session["uid"])).fetchone()
    if row:
        for a in d.execute("SELECT url, kind FROM hw_attach WHERE hw_id=?", (hid,)):
            if a["kind"] == "file":
                fp = os.path.join(UPLOADS, a["url"])
                if os.path.isfile(fp):
                    try: os.remove(fp)
                    except OSError: pass
        d.execute("DELETE FROM homework WHERE id=?", (hid,))
        d.commit()
    return redirect(url_for("homework"))

@app.post("/homework/done")
@role("student")
def homework_done():
    j, d, uid = request.get_json(), db(), session["uid"]
    if not d.execute("SELECT 1 FROM homework h JOIN users u ON u.group_id=h.group_id WHERE h.id=? AND u.id=?",
                     (j["id"], uid)).fetchone():
        abort(403)
    if j["done"]:
        d.execute("INSERT OR IGNORE INTO hw_done VALUES(?,?)", (j["id"], uid))
    else:
        d.execute("DELETE FROM hw_done WHERE hw_id=? AND student_id=?", (j["id"], uid))
    d.commit()
    return jsonify(ok=True)

@app.post("/homework/note")
@role("student")
def homework_note():
    j, d, uid = request.get_json(), db(), session["uid"]
    if not d.execute("SELECT 1 FROM homework h JOIN users u ON u.group_id=h.group_id WHERE h.id=? AND u.id=?",
                     (j["id"], uid)).fetchone():
        abort(403)
    body = (j.get("body") or "").strip()
    now = datetime.now().strftime("%d.%m.%Y %H:%M")
    if body:
        d.execute("""INSERT INTO hw_notes(hw_id,student_id,body,updated) VALUES(?,?,?,?)
          ON CONFLICT(hw_id,student_id) DO UPDATE SET body=excluded.body, updated=excluded.updated""",
                  (j["id"], uid, body, now))
    else:
        d.execute("DELETE FROM hw_notes WHERE hw_id=? AND student_id=?", (j["id"], uid))
    d.commit()
    return jsonify(ok=True, updated=now)

@app.route("/uploads/<path:name>")
def serve_upload(name):
    if "uid" not in session:
        abort(403)
    # only basename
    name = os.path.basename(name)
    path = os.path.join(UPLOADS, name)
    if not os.path.isfile(path):
        abort(404)
    # check access: teacher of assignment or student of group
    att = db().execute("SELECT a.hw_id, h.group_id, h.subject_id FROM hw_attach a JOIN homework h ON h.id=a.hw_id WHERE a.url=? AND a.kind='file'",
                       (name,)).fetchone()
    if not att:
        abort(404)
    role, uid = session["role"], session["uid"]
    d = db()
    if role == "teacher":
        if not d.execute("SELECT 1 FROM assignments WHERE teacher_id=? AND group_id=? AND subject_id=?",
                         (uid, att["group_id"], att["subject_id"])).fetchone():
            abort(403)
    elif role == "student":
        if not d.execute("SELECT 1 FROM users WHERE id=? AND group_id=?", (uid, att["group_id"])).fetchone():
            abort(403)
    else:
        abort(403)
    from flask import send_from_directory
    return send_from_directory(UPLOADS, name, as_attachment=False)


if __name__ == "__main__":
    setup()
    app.run(debug=True)
