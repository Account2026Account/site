# -*- coding: utf-8 -*-
from flask import render_template, request, redirect, url_for, session, jsonify, flash, abort, Response, send_from_directory
from datetime import date, datetime, timedelta
from routes import FlatBlueprint
from core import db, role, stats, has_journal_access, grant_journal_access, notify, notify_many, setting, minutes, gap_label, overlap, occurrences, sync_lessons, feed, Bad, WD, WDF, UPLOADS, ALLOWED_EXT, _ROOT, DB, migrate, seed_marks, setup, can_fill_attendance
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename
import io, re, os, sqlite3, shutil, tempfile, random

bp = FlatBlueprint("schedule", __name__)
from core import bell_rows, scope

# ---------- Расписание ----------
@bp.route("/schedule")
def schedule():
    if "uid" not in session: return redirect(url_for("login"))
    d, off = db(), request.args.get("week", 0, type=int)
    gid, tid, groups = scope(d)
    mon = date.today() - timedelta(days=date.today().weekday()) + timedelta(weeks=off)
    occ = occurrences(d, mon, mon + timedelta(days=5), gid=gid, tid=tid)
    days = []
    for i in range(6):
        day = mon + timedelta(days=i); lst = [o for o in occ if o["day"] == day.isoformat()]
        if gid or tid:
            for x, y in zip(lst, lst[1:]):
                g = minutes(y["start"]) - minutes(x["end"])
                if g > 0: y["gap_label"] = gap_label(g)
        days.append(dict(date=day, lst=lst))
    is_teacher = session.get("role") == "teacher"
    return render_template("schedule.html", days=days, off=off, gid=gid if gid and gid > 0 else None, groups=groups, wd=WDF,
                           today=date.today(), bells=bell_rows(d.execute("SELECT * FROM bells ORDER BY num").fetchall()),
                           is_teacher=is_teacher, show_group_filter=bool(groups))

@bp.route("/schedule.xlsx")
def schedule_xlsx():
    """Экспорт расписания на неделю в Excel (сетка + список) с оформлением."""
    if "uid" not in session: return redirect(url_for("login"))
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter as L
    d, off = db(), request.args.get("week", 0, type=int)
    gid, tid, groups = scope(d)
    mon = date.today() - timedelta(days=date.today().weekday()) + timedelta(weeks=off)
    occ = occurrences(d, mon, mon + timedelta(days=5), gid=gid, tid=tid)
    bells = d.execute("SELECT * FROM bells ORDER BY num").fetchall()

    wb = Workbook()
    ws = wb.active
    ws.title = "Расписание"
    ws.sheet_view.showGridLines = False
    fill = lambda h: PatternFill("solid", fgColor=h)
    thin = Side(style="thin", color="D9DEEC")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="top", wrap_text=True)

    title = "Расписание"
    if gid and groups:
        gname = next((g["name"] for g in groups if g["id"] == gid), None)
        if gname:
            title += f" · {gname}"
    elif session.get("role") == "teacher":
        title += " · мои занятия"
    title += f"  {mon.strftime('%d.%m.%Y')} – {(mon + timedelta(days=5)).strftime('%d.%m.%Y')}"

    ws["A1"] = title
    ws["A1"].font = Font(size=16, bold=True, color="14213D")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=7)
    ws["A2"] = f"Выгружено {datetime.now():%d.%m.%Y %H:%M} · Campus"
    ws["A2"].font = Font(color="66708C", size=10)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=7)

    # Заголовки
    ws.cell(4, 1, "Пара / время").font = Font(bold=True, color="FFFFFF", size=11)
    ws.cell(4, 1).fill = fill("14213D")
    ws.cell(4, 1).alignment = center
    ws.cell(4, 1).border = box
    for i in range(6):
        cell = ws.cell(4, 2 + i, f"{WDF[i]}\n{(mon + timedelta(days=i)).strftime('%d.%m')}")
        cell.font = Font(bold=True, color="FFFFFF", size=11)
        cell.fill = fill("14213D")
        cell.alignment = center
        cell.border = box
    ws.row_dimensions[4].height = 36

    by_day_bell = {}
    for o in occ:
        by_day_bell.setdefault((o["day"], o["num"]), []).append(o)

    for ri, b in enumerate(bells, 5):
        time_cell = ws.cell(ri, 1, f'{b["num"]} пара\n{b["start"]}–{b["end"]}')
        time_cell.font = Font(bold=True, size=10)
        time_cell.alignment = center
        time_cell.border = box
        time_cell.fill = fill("EEF1F8")
        max_lines = 1
        for di in range(6):
            day = (mon + timedelta(days=di)).isoformat()
            items = by_day_bell.get((day, b["num"]), [])
            cell = ws.cell(ri, 2 + di)
            cell.border = box
            cell.alignment = left
            if not items:
                cell.value = ""
                continue
            blocks = []
            for o in items:
                st = ""
                if o["status"] == "cancel":
                    st = " [отменено]"
                elif o["status"] == "replace":
                    st = " [замена]"
                block = f'{o["sname"]}{st}\n{o["gname"]}\n{o["tname"]}\nауд. {o["room"] or "—"}'
                if o.get("note"):
                    block += f'\n{o["note"]}'
                blocks.append(block)
            text = "\n———\n".join(blocks)
            cell.value = text
            max_lines = max(max_lines, text.count("\n") + 1)
            if any(o["status"] == "cancel" for o in items):
                cell.fill = fill("FFE4E4")
            elif any(o["status"] == "replace" for o in items):
                cell.fill = fill("FFFBE8")
            else:
                cell.fill = fill("F5F7FE")
        ws.row_dimensions[ri].height = max(50, 14 * max_lines + 8)

    ws.column_dimensions["A"].width = 14
    for i in range(2, 8):
        ws.column_dimensions[L(i)].width = 22
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.freeze_panes = "B5"

    # Лист «Список»
    ws2 = wb.create_sheet("Список")
    ws2.sheet_view.showGridLines = False
    ws2["A1"] = title
    ws2["A1"].font = Font(size=14, bold=True, color="14213D")
    ws2.merge_cells("A1:J1")
    cols = ["Дата", "День", "Пара", "Время", "Предмет", "Группа", "Преподаватель", "Аудитория", "Статус", "Примечание"]
    for i, h in enumerate(cols, 1):
        x = ws2.cell(3, i, h)
        x.font = Font(bold=True, color="FFFFFF")
        x.fill = fill("14213D")
        x.alignment = center
        x.border = box
    for ri, o in enumerate(sorted(occ, key=lambda x: (x["day"], x["start"])), 4):
        st = {"ok": "ок", "cancel": "отменено", "replace": "замена"}.get(o["status"], o["status"])
        vals = [
            f'{o["day"][8:10]}.{o["day"][5:7]}.{o["day"][:4]}',
            WDF[o["weekday"]],
            o["num"],
            f'{o["start"]}–{o["end"]}',
            o["sname"],
            o["gname"],
            o["tname"],
            o["room"] or "—",
            st,
            o.get("note") or "",
        ]
        for ci, v in enumerate(vals, 1):
            x = ws2.cell(ri, ci, v)
            x.border = box
            x.alignment = center if ci <= 4 else left
            if o["status"] == "cancel":
                x.fill = fill("FFE4E4")
            elif o["status"] == "replace":
                x.fill = fill("FFFBE8")
    for i, w in enumerate([12, 10, 8, 12, 22, 12, 20, 12, 12, 24], 1):
        ws2.column_dimensions[L(i)].width = w
    if occ:
        ws2.auto_filter.ref = f"A3:J{3 + len(occ)}"
    ws2.freeze_panes = "A4"

    buf = io.BytesIO()
    wb.save(buf)
    fname = f"schedule_{mon.strftime('%Y%m%d')}.xlsx"
    return Response(
        buf.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={fname}"},
    )


# ---------- Импорт расписания из файлов ----------
WD_RU = {"понедельник": 0, "вторник": 1, "среда": 2, "четверг": 3, "пятница": 4, "суббота": 5}
WD_SHORT = {"пн": 0, "вт": 1, "ср": 2, "чт": 3, "пт": 4, "сб": 5}


# --- import parsers ---
try:
    from services.importers import (
        import_weekly_docx, import_replacements_pdf, import_rooms_pdf,
        cleanup_bogus_bells,
    )
except ImportError:
    import sys
    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
    from services.importers import (
        import_weekly_docx, import_replacements_pdf, import_rooms_pdf,
        cleanup_bogus_bells,
    )

@bp.route("/dispatcher/import", methods=["GET", "POST"])
@role("dispatcher", "admin")
def dispatcher_import():
    if request.method == "GET":
        return redirect(url_for("dispatcher"))
    d = db()
    f = request.files.get("file")
    kind = (request.form.get("kind") or "auto").strip()
    day = (request.form.get("day") or "").strip() or None
    if day:
        try:
            day = date.fromisoformat(day).isoformat()
        except Exception:
            day = None
    if not f or not f.filename:
        flash("Выберите файл")
        return redirect(url_for("dispatcher"))
    fname = secure_filename(f.filename)
    ext = fname.rsplit(".", 1)[-1].lower() if "." in fname else ""
    tmp = tempfile.NamedTemporaryFile(delete=False, suffix="." + ext)
    try:
        f.save(tmp.name)
        tmp.close()
        if kind == "auto":
            if ext in ("docx",):
                kind = "weekly"
            elif any(x in fname.lower() for x in ("kabinet", "кабинет", "auditori", "аудитор", "rooms")):
                kind = "rooms"
            elif any(x in fname.lower() for x in ("zamen", "замен", "изменен", "change")):
                kind = "replace"
            elif ext == "pdf":
                # без подсказки в имени — спрашиваем через kind в форме; по умолчанию замены
                kind = "replace"
            else:
                flash("Неизвестный тип файла. Используйте DOCX (неделя) или PDF (замены/кабинеты).")
                return redirect(url_for("dispatcher"))
        if kind == "weekly":
            if ext != "docx":
                flash("Недельное расписание: нужен файл .docx")
                return redirect(url_for("dispatcher"))
            st = import_weekly_docx(d, tmp.name)
            try:
                cleanup_bogus_bells(d)
            except Exception:
                pass
            flash(f"Импорт недели: групп {st['groups']}, занятий {st['lessons']}. "
                  f"Новые группы: {len(st['created_groups'])}, предметы: {len(set(st['created_subjects']))}.")
            try:
                sync_lessons(d)
            except Exception:
                pass
        elif kind == "replace":
            st = import_replacements_pdf(d, tmp.name, day)
            flash(f"Импорт замен на {st.get('day') or '?'}: замен {st['replacements']}, отмен {st['cancels']}.")
            try:
                sync_lessons(d)
            except Exception:
                pass
        elif kind == "rooms":
            st = import_rooms_pdf(d, tmp.name, day)
            flash(f"Импорт кабинетов на {st.get('day') or '?'}: кабинетов {st['rooms']}, отмен {st.get('cancels',0)}, "
                  f"групп в файле {st.get('groups_seen',0)}, новых групп {st.get('groups_created',0)}. "
                  f"{'Сначала загрузите недельное расписание — тогда кабинеты привяжутся к парам.' if st['rooms']==0 and st.get('groups_created',0)>=0 else ''}")
        else:
            flash("Неизвестный режим импорта")
    except Exception as e:
        d.rollback()
        flash(f"Ошибка импорта: {e}")
    finally:
        try:
            os.unlink(tmp.name)
        except Exception:
            pass
    return redirect(url_for("dispatcher"))


@bp.route("/dispatcher")
@role("dispatcher", "admin")
def dispatcher():
    d = db()
    groups = d.execute("SELECT * FROM groups ORDER BY name").fetchall()
    gid = request.args.get("group", type=int) or (groups[0]["id"] if groups else None)
    view = request.args.get("view") or "week"  # kept for compatibility, UI now shows both
    off = request.args.get("week", 0, type=int)
    mon = date.today() - timedelta(days=date.today().weekday()) + timedelta(weeks=off)
    grid, once, occ, wgrid = {}, [], [], {}
    if gid:
        for e in d.execute("""SELECT s.*, sub.name sname, t.name tname FROM schedule s JOIN subjects sub ON sub.id=s.subject_id
                JOIN users t ON t.id=s.teacher_id WHERE s.group_id=? ORDER BY s.parity""", (gid,)):
            e = dict(e)
            extra = (e.get("teachers_extra") or "").strip()
            if extra:
                extras = [x.strip() for x in extra.split(",") if x.strip() and x.strip().lower() != (e["tname"] or "").lower()]
                if extras:
                    e["tname"] = e["tname"] + ", " + ", ".join(extras)
            if e["day"]: once.append(e)
            else: grid.setdefault((e["weekday"], e["bell_id"]), []).append(e)
        once.sort(key=lambda e: e["day"] or "")
        occ = occurrences(d, mon, mon + timedelta(days=5), gid=gid)
        for o in occ: wgrid.setdefault((o["weekday"], o["num"]), []).append(o)
    return render_template("dispatcher.html", groups=groups, gid=gid, view=view, off=off, grid=grid, once=once, occ=occ, wgrid=wgrid,
        dates=[mon + timedelta(days=i) for i in range(6)], wd=WD, term=setting(d, "term_start"),
        bells=bell_rows(d.execute("SELECT * FROM bells ORDER BY num").fetchall()),
        subjects=d.execute("SELECT * FROM subjects ORDER BY name").fetchall(),
        teachers=d.execute("SELECT id, name FROM users WHERE role='teacher' ORDER BY name").fetchall())

@bp.post("/dispatcher/<act>")
@role("dispatcher", "admin")
def dispatcher_act(act):
    d, f = db(), request.form
    gid = f.get("group") or None
    try:
        if act == "bell":
            d.execute("INSERT INTO bells(num,start,end) VALUES(?,?,?)", (int(f["num"]), f["start"], f["end"]))
        elif act == "bell_del":
            d.execute("DELETE FROM bells WHERE id=?", (f["id"],))
        elif act == "term":
            d.execute("INSERT OR REPLACE INTO settings VALUES('term_start',?)", (date.fromisoformat(f["value"]).isoformat(),))
        elif act == "entry":
            new = dict(group_id=int(f["group"]), subject_id=int(f["subject_id"]), teacher_id=int(f["teacher_id"]), bell_id=int(f["bell_id"]),
                       weekday=int(f["weekday"]), day=f.get("day") or None, parity=int(f["parity"]))
            if new["day"]:
                new["day"] = date.fromisoformat(new["day"]).isoformat(); new["weekday"] = None
            for e in d.execute("SELECT * FROM schedule WHERE bell_id=?", (new["bell_id"],)):
                if (e["group_id"] == new["group_id"] or e["teacher_id"] == new["teacher_id"]) and overlap(e, new):
                    raise Bad("В это время у группы или у преподавателя уже есть занятие")
            d.execute("INSERT INTO schedule(group_id,subject_id,teacher_id,bell_id,weekday,day,room,parity) VALUES(?,?,?,?,?,?,?,?)",
                      (new["group_id"], new["subject_id"], new["teacher_id"], new["bell_id"], new["weekday"], new["day"], f.get("room", "").strip(), new["parity"]))
            d.execute("INSERT OR IGNORE INTO assignments(teacher_id,group_id,subject_id) VALUES(?,?,?)", (new["teacher_id"], new["group_id"], new["subject_id"]))
        elif act == "entry_del":
            d.execute("DELETE FROM schedule WHERE id=?", (f["id"],))
        elif act == "change":
            s = d.execute("SELECT * FROM schedule WHERE id=?", (f["schedule_id"],)).fetchone()
            if not s: raise Bad("Занятие не найдено")
            kind, day, room = f["kind"], date.fromisoformat(f["day"]).isoformat(), f.get("room", "").strip()
            tid = int(f["teacher_id"]) if f.get("teacher_id") else None
            if kind == "replace" and not (tid or room): raise Bad("Для замены укажите преподавателя или аудиторию")
            rep_ = kind == "replace"
            d.execute("""INSERT INTO changes(schedule_id,day,kind,teacher_id,room,note) VALUES(?,?,?,?,?,?)
              ON CONFLICT(schedule_id,day) DO UPDATE SET kind=excluded.kind, teacher_id=excluded.teacher_id, room=excluded.room, note=excluded.note""",
                      (s["id"], day, kind, tid if rep_ else None, room if rep_ else "", f.get("note", "").strip()))
            if kind == "cancel":  # пустое занятие убираем из журнала, с отметками оставляем
                d.execute("""DELETE FROM lessons WHERE group_id=? AND subject_id=? AND day=?
                  AND NOT EXISTS(SELECT 1 FROM marks WHERE lesson_id=lessons.id)""", (s["group_id"], s["subject_id"], day))
            elif tid:  # заменяющий преподаватель получает доступ к журналу группы
                d.execute("INSERT OR IGNORE INTO assignments(teacher_id,group_id,subject_id) VALUES(?,?,?)", (tid, s["group_id"], s["subject_id"]))
        elif act == "change_del":
            d.execute("DELETE FROM changes WHERE id=?", (f["id"],))
        elif act == "reset_schedule":
            # scope: all | group ; what: weekly | day | changes | all
            scope = (f.get("scope") or "group").strip()
            what = (f.get("what") or "all").strip()
            day = (f.get("day") or "").strip() or None
            if day:
                try:
                    day = date.fromisoformat(day).isoformat()
                except Exception:
                    day = None
            target_gid = int(f["group"]) if f.get("group") else None
            if scope == "group" and not target_gid:
                raise Bad("Выберите группу для сброса")
            if what == "day" and not day:
                raise Bad("Укажите дату для сброса замен/разовых на день")

            def _sch_filter(sql_extra="", params=()):
                if scope == "group":
                    return sql_extra + (" AND group_id=?" if "group_id" not in sql_extra else ""), params + (target_gid,)
                return sql_extra, params

            if what in ("weekly", "all"):
                # постоянное расписание (weekday, без day)
                if scope == "group":
                    d.execute("DELETE FROM schedule WHERE group_id=? AND weekday IS NOT NULL AND (day IS NULL OR day='')", (target_gid,))
                else:
                    d.execute("DELETE FROM schedule WHERE weekday IS NOT NULL AND (day IS NULL OR day='')")
            if what in ("day", "all"):
                if day:
                    # разовые на день + changes на день
                    if scope == "group":
                        ids = [r["id"] for r in d.execute("SELECT id FROM schedule WHERE group_id=?", (target_gid,))]
                        if ids:
                            q = ",".join("?" * len(ids))
                            d.execute(f"DELETE FROM changes WHERE day=? AND schedule_id IN ({q})", [day] + ids)
                        d.execute("DELETE FROM schedule WHERE group_id=? AND day=?", (target_gid, day))
                    else:
                        d.execute("DELETE FROM changes WHERE day=?", (day,))
                        d.execute("DELETE FROM schedule WHERE day=?", (day,))
                elif what == "all":
                    # все разовые и все замены
                    if scope == "group":
                        ids = [r["id"] for r in d.execute("SELECT id FROM schedule WHERE group_id=?", (target_gid,))]
                        if ids:
                            q = ",".join("?" * len(ids))
                            d.execute(f"DELETE FROM changes WHERE schedule_id IN ({q})", ids)
                        d.execute("DELETE FROM schedule WHERE group_id=? AND day IS NOT NULL AND day!=''", (target_gid,))
                    else:
                        d.execute("DELETE FROM changes")
                        d.execute("DELETE FROM schedule WHERE day IS NOT NULL AND day!=''")
            if what == "changes":
                if day:
                    if scope == "group":
                        ids = [r["id"] for r in d.execute("SELECT id FROM schedule WHERE group_id=?", (target_gid,))]
                        if ids:
                            q = ",".join("?" * len(ids))
                            d.execute(f"DELETE FROM changes WHERE day=? AND schedule_id IN ({q})", [day] + ids)
                    else:
                        d.execute("DELETE FROM changes WHERE day=?", (day,))
                else:
                    if scope == "group":
                        ids = [r["id"] for r in d.execute("SELECT id FROM schedule WHERE group_id=?", (target_gid,))]
                        if ids:
                            q = ",".join("?" * len(ids))
                            d.execute(f"DELETE FROM changes WHERE schedule_id IN ({q})", ids)
                    else:
                        d.execute("DELETE FROM changes")
            if what == "all" and not day:
                # полный сброс слотов
                if scope == "group":
                    d.execute("DELETE FROM schedule WHERE group_id=?", (target_gid,))
                else:
                    d.execute("DELETE FROM schedule")
                    d.execute("DELETE FROM changes")
            flash("Расписание сброшено")
            d.commit()
            return redirect(url_for("dispatcher", group=gid, view=f.get("view") or None, week=f.get("week") or None))
        d.commit(); flash("Готово")
    except Bad as e:
        d.rollback(); flash(str(e))
    except sqlite3.IntegrityError:
        d.rollback(); flash("Такая запись уже существует")
    except (ValueError, KeyError):
        d.rollback(); flash("Проверьте введённые данные")
    return redirect(url_for("dispatcher", group=gid, view=f.get("view") or None, week=f.get("week") or None))

