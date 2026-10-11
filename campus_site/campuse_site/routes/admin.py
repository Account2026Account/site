# -*- coding: utf-8 -*-
from flask import render_template, request, redirect, url_for, session, jsonify, flash, abort, Response, send_from_directory
from datetime import date, datetime, timedelta
from routes import FlatBlueprint
from core import db, role, stats, has_journal_access, grant_journal_access, notify, notify_many, setting, minutes, gap_label, overlap, occurrences, sync_lessons, feed, Bad, WD, WDF, UPLOADS, ALLOWED_EXT, _ROOT, DB, migrate, seed_marks, setup, can_fill_attendance
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename
import io, re, os, sqlite3, shutil, tempfile, random

bp = FlatBlueprint("admin", __name__)

# ---------- Администратор ----------
@bp.route("/admin")
@role("admin")
def admin():
    d = db(); q = lambda x: d.execute(x).fetchall(); one = lambda x: d.execute(x).fetchone()[0]
    users = q("""SELECT u.*, g.name gname FROM users u LEFT JOIN groups g ON g.id=u.group_id
      ORDER BY CASE u.role WHEN 'admin' THEN 0 WHEN 'dispatcher' THEN 1 WHEN 'teacher' THEN 2 ELSE 3 END, g.name, u.name""")
    sbg = {}
    for u in users:
        if u["role"] == "student" and u["group_id"]: sbg.setdefault(u["group_id"], []).append(u)
    jacc = q("""SELECT ja.teacher_id, ja.group_id, ja.subject_id, ja.level, t.name tname, g.name gname, s.name sname
                FROM journal_access ja
                JOIN users t ON t.id=ja.teacher_id JOIN groups g ON g.id=ja.group_id JOIN subjects s ON s.id=ja.subject_id
                ORDER BY g.name, s.name, t.name""")
    students_list = [u for u in users if u["role"] == "student"]
    staff_list = [u for u in users if u["role"] in ("teacher", "admin", "dispatcher", "curator")]
    return render_template("admin.html", users=users, sbg=sbg, groups=q("SELECT * FROM groups ORDER BY name"),
        subjects=q("SELECT * FROM subjects ORDER BY name"), teachers=q("SELECT * FROM users WHERE role='teacher' ORDER BY name"),
        assigns=q("""SELECT a.id, g.name gname, s.name sname, t.name tname FROM assignments a JOIN groups g ON g.id=a.group_id
                     JOIN subjects s ON s.id=a.subject_id JOIN users t ON t.id=a.teacher_id ORDER BY g.name, s.name, t.name"""),
        journal_access=jacc,
        counts=dict(groups=one("SELECT COUNT(*) FROM groups"), students=one("SELECT COUNT(*) FROM users WHERE role='student'"),
                    teachers=one("SELECT COUNT(*) FROM users WHERE role='teacher'"), lessons=one("SELECT COUNT(*) FROM lessons")),
        anns=announcements(), can_del=True, all_ok=True, students=students_list, staff=staff_list)

@bp.post("/admin/<act>")
@role("admin")
def admin_act(act):
    d, f = db(), request.form
    try:
        if act == "group":
            d.execute("INSERT INTO groups(name) VALUES(?)", (f["name"].strip(),))
        elif act == "subject":
            d.execute("INSERT INTO subjects(name) VALUES(?)", (f["name"].strip(),))
        elif act == "user":
            r, gid = f["role"], f.get("group_id") or None
            if r not in ("teacher", "student", "dispatcher", "curator"): raise ValueError("Выберите роль")
            if r == "student" and not gid: raise ValueError("Для студента выберите группу")
            d.execute("INSERT INTO users(login,pw,name,role,group_id) VALUES(?,?,?,?,?)",
                      (f["login"].strip(), generate_password_hash(f["password"]), f["name"].strip(), r, gid if r == "student" else None))
            if r == "curator" and gid:
                uid = d.execute("SELECT id FROM users WHERE login=?", (f["login"].strip(),)).fetchone()[0]
                d.execute("UPDATE groups SET curator_id=? WHERE id=?", (uid, gid))
        elif act == "assign":
            d.execute("INSERT INTO assignments(teacher_id,group_id,subject_id) VALUES(?,?,?)", (f["teacher_id"], f["group_id"], f["subject_id"]))
            grant_journal_access(int(f["teacher_id"]), int(f["group_id"]), int(f["subject_id"]), "edit", session["uid"])
        elif act == "journal_access":
            tid, gid, sid = int(f["teacher_id"]), int(f["group_id"]), int(f["subject_id"])
            level = f.get("level") or "edit"
            if level not in ("edit", "view"):
                raise ValueError("Уровень: edit или view")
            grant_journal_access(tid, gid, sid, level, session["uid"])
        elif act == "journal_access_del":
            d.execute("DELETE FROM journal_access WHERE teacher_id=? AND group_id=? AND subject_id=?",
                      (f["teacher_id"], f["group_id"], f["subject_id"]))
        elif act == "head":
            sid = f.get("student_id") or None
            if sid and not d.execute("SELECT 1 FROM users WHERE id=? AND group_id=? AND role='student'", (sid, f["group_id"])).fetchone():
                raise ValueError("Староста должен быть студентом этой группы")
            d.execute("UPDATE groups SET head_id=? WHERE id=?", (sid, f["group_id"]))
        elif act == "deputy_head":
            sid = f.get("student_id") or None
            if sid and not d.execute("SELECT 1 FROM users WHERE id=? AND group_id=? AND role='student'", (sid, f["group_id"])).fetchone():
                raise ValueError("Зам. старосты должен быть студентом этой группы")
            d.execute("UPDATE groups SET deputy_head_id=? WHERE id=?", (sid, f["group_id"]))
        elif act == "curator":
            tid = f.get("teacher_id") or None
            if tid and not d.execute("SELECT 1 FROM users WHERE id=? AND role IN ('teacher','curator','admin')", (tid,)).fetchone():
                raise ValueError("Куратор должен быть преподавателем или куратором")
            d.execute("UPDATE groups SET curator_id=? WHERE id=?", (tid or None, f["group_id"]))
            # доступ куратора — через groups.curator_id в has_journal_access, без journal_access
        elif act == "student_info":
            # расширенные данные студента
            uid = int(f["id"])
            d.execute("""UPDATE users SET email=?, phone=?, birth_date=?, student_id_number=?,
                         parent_name=?, parent_phone=?, address=?, notes=? WHERE id=? AND role='student'""",
                      (f.get("email","").strip(), f.get("phone","").strip(), f.get("birth_date","").strip(),
                       f.get("student_id_number","").strip(), f.get("parent_name","").strip(),
                       f.get("parent_phone","").strip(), f.get("address","").strip(),
                       f.get("notes","").strip(), uid))
        elif act == "toggle_active":
            uid = int(f["id"])
            d.execute("UPDATE users SET active=CASE WHEN COALESCE(active,1)=1 THEN 0 ELSE 1 END WHERE id=? AND id!=?",
                      (uid, session["uid"]))
        elif act == "move":
            uid_move = int(f["id"])
            new_gid = f.get("group_id") or None
            # снять со старосты/зама, если уходит из группы (или меняет группу)
            old = d.execute("SELECT group_id FROM users WHERE id=? AND role='student'", (uid_move,)).fetchone()
            old_gid = old["group_id"] if old else None
            if old_gid and str(old_gid) != str(new_gid or ""):
                d.execute("UPDATE groups SET head_id=NULL WHERE id=? AND head_id=?", (old_gid, uid_move))
                d.execute("UPDATE groups SET deputy_head_id=NULL WHERE id=? AND deputy_head_id=?", (old_gid, uid_move))
            d.execute("UPDATE users SET group_id=? WHERE id=? AND role='student'", (new_gid, uid_move))
        elif act == "pw":
            if len(f["password"]) < 6: raise ValueError("Пароль должен быть не короче 6 символов")
            d.execute("UPDATE users SET pw=? WHERE id=?", (generate_password_hash(f["password"]), f["id"]))
        elif act == "del":
            t = {"group": "groups", "subject": "subjects", "user": "users", "assign": "assignments"}[f["kind"]]
            if t == "users" and int(f["id"]) == session["uid"]: raise ValueError("Нельзя удалить самого себя")
            if t == "groups":  # студенты остаются в системе, но без группы
                d.execute("UPDATE users SET group_id=NULL WHERE group_id=?", (f["id"],))
            d.execute(f"DELETE FROM {t} WHERE id=?", (f["id"],))
        d.commit(); flash("Готово")
    except sqlite3.IntegrityError:
        d.rollback(); flash("Не получилось: такая запись уже есть, либо на неё ссылаются (например, в группе остались студенты).")
    except (ValueError, KeyError) as e:
        d.rollback(); flash(str(e))
    return redirect(url_for("admin"))


# ---------- Глобальная выгрузка ----------
# ---------- Глобальная выгрузка оценок / посещаемости ----------
@bp.route("/export/student/<int:uid>")
@role("admin", "teacher")
def export_student(uid):
    """Подробная выгрузка студента: сводка, все предметы с оценками/посещаемостью/итогами/ДЗ."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter as L
    d = db()
    st = d.execute("SELECT u.*, g.name gname FROM users u LEFT JOIN groups g ON g.id=u.group_id WHERE u.id=? AND u.role='student'", (uid,)).fetchone()
    if not st:
        abort(404)
    if session.get("role") == "teacher":
        ok = d.execute("SELECT 1 FROM assignments a WHERE a.teacher_id=? AND a.group_id=?",
                       (session["uid"], st["group_id"])).fetchone()
        if not ok:
            abort(403)
    wb = Workbook()
    fill = lambda h: PatternFill("solid", fgColor=h)
    thin = Side(style="thin", color="D9DEEC"); box = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    gfill = {2: "FFC9C9", 3: "FFD84D", 4: "CFE0FF", 5: "B8F0D3"}
    head_font = Font(bold=True, color="FFFFFF")
    title_font = Font(size=14, bold=True, color="14213D")

    subjects = []
    if st["group_id"]:
        subjects = d.execute("""SELECT s.id, s.name,
            (SELECT group_concat(t.name, ', ') FROM assignments a2 JOIN users t ON t.id=a2.teacher_id
              WHERE a2.group_id=? AND a2.subject_id=s.id) tname
            FROM subjects s WHERE s.id IN (SELECT subject_id FROM assignments WHERE group_id=?)
            ORDER BY s.name""", (st["group_id"], st["group_id"])).fetchall()
        # also subjects that have marks even without assignment
        extra = d.execute("""SELECT DISTINCT s.id, s.name FROM subjects s
            JOIN lessons l ON l.subject_id=s.id JOIN marks m ON m.lesson_id=l.id
            WHERE m.student_id=? AND l.group_id=?""", (uid, st["group_id"])).fetchall()
        known = {s["id"] for s in subjects}
        for e in extra:
            if e["id"] not in known:
                subjects.append(e)
                known.add(e["id"])

    # ===== Sheet 1: Сводка =====
    ws = wb.active; ws.title = "Сводка"
    ws["A1"] = f'Студент: {st["name"]}'; ws["A1"].font = title_font
    ws["A2"] = f'Группа: {st["gname"] or "—"}. Логин: {st["login"]}. Выгружено {datetime.now():%d.%m.%Y %H:%M}'
    ws["A2"].font = Font(color="66708C")
    heads = ["№", "Предмет", "Преподаватели", "Ср. балл", "Оценок", "Был", "Не был", "Уваж.", "Пропуски",
             "Посещ. %", "Доп. баллы", "Итоги"]
    for i, h in enumerate(heads, 1):
        x = ws.cell(4, i, h); x.font = head_font; x.fill = fill("14213D"); x.border = box; x.alignment = center
    row_i = 5
    all_cells = []
    for si, sub in enumerate(subjects, 1):
        cells = d.execute("""SELECT m.present, m.grade, m.excuse FROM lessons l
            LEFT JOIN marks m ON m.lesson_id=l.id AND m.student_id=?
            WHERE l.group_id=? AND l.subject_id=?""", (uid, st["group_id"], sub["id"])).fetchall()
        all_cells.extend(cells)
        stt = stats(cells)
        grades_n = sum(1 for c in cells if c and c["grade"])
        was = sum(1 for c in cells if c and c["present"] == 1)
        bp = d.execute("SELECT balance FROM bonus_points WHERE student_id=? AND subject_id=?", (uid, sub["id"])).fetchone()
        kinds = get_final_kinds(d, st["group_id"], sub["id"])
        frows = d.execute("SELECT kind, grade FROM final_grades WHERE student_id=? AND subject_id=?", (uid, sub["id"])).fetchall()
        by_k = {r["kind"]: r["grade"] for r in frows if r["grade"]}
        finals_txt = "; ".join(f"{kinds.get(k,k)}: {g}" for k, g in by_k.items()) or "—"
        tname = sub["tname"] if "tname" in sub.keys() else ""
        vals = [si, sub["name"], tname or "—", stt["avg"], grades_n, was,
                stt.get("unexcused") or 0, stt.get("excused") or 0, stt["missed"],
                stt["pct"], bp["balance"] if bp else 0, finals_txt]
        for i, v in enumerate(vals, 1):
            x = ws.cell(row_i, i, v); x.border = box
            x.alignment = center if i not in (2, 3, 12) else left
            if i == 4 and v is not None and v < 3:
                x.fill = fill("FFE4E4")
            if i == 10 and v is not None and v < 70:
                x.fill = fill("FFF3CD")
        row_i += 1
    # totals
    tot = stats(all_cells)
    ws.cell(row_i + 1, 1, "ИТОГО").font = Font(bold=True)
    ws.cell(row_i + 1, 4, tot["avg"]); ws.cell(row_i + 1, 9, tot["missed"]); ws.cell(row_i + 1, 10, tot["pct"])
    for i in range(1, 13):
        ws.cell(row_i + 1, i).border = box; ws.cell(row_i + 1, i).font = Font(bold=True)
        ws.cell(row_i + 1, i).fill = fill("EEF1F8")
    widths = [5, 28, 28, 11, 10, 8, 10, 10, 10, 11, 12, 40]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[L(i)].width = w
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:L{max(4, row_i-1)}"

    # ===== Sheet 2: Все оценки (плоская таблица) =====
    wsA = wb.create_sheet("Все оценки")
    wsA["A1"] = f'Все отметки: {st["name"]} ({st["gname"] or "—"})'
    wsA["A1"].font = title_font
    ah = ["Предмет", "Дата", "Тема", "Пара", "Статус", "Оценка", "Причина уваж.", "Доп. баллы"]
    for i, h in enumerate(ah, 1):
        x = wsA.cell(3, i, h); x.font = head_font; x.fill = fill("14213D"); x.border = box; x.alignment = center
    ri = 4
    for sub in subjects:
        lessons = d.execute("""SELECT l.id, l.day, l.topic, m.present, m.grade, m.excuse,
            (SELECT points_used FROM bonus_applications ba WHERE ba.lesson_id=l.id AND ba.student_id=? AND ba.cancelled_at IS NULL) pts
            FROM lessons l LEFT JOIN marks m ON m.lesson_id=l.id AND m.student_id=?
            WHERE l.group_id=? AND l.subject_id=? ORDER BY l.day""", (uid, uid, st["group_id"], sub["id"])).fetchall()
        for l in lessons:
            status = "—"
            if l["grade"]:
                status = "оценка"
            elif l["present"] == 2:
                status = "уваж."
            elif l["present"] == 0:
                status = "н"
            elif l["present"] == 1:
                status = "был"
            vals = [sub["name"],
                    f'{l["day"][8:10]}.{l["day"][5:7]}.{l["day"][:4]}' if l["day"] else "",
                    l["topic"] or "", "", status, l["grade"] or "",
                    (l["excuse"] or "") if l["present"] == 2 else "",
                    l["pts"] or ""]
            for i, v in enumerate(vals, 1):
                x = wsA.cell(ri, i, v); x.border = box
                if l["grade"] and l["grade"] in gfill and i == 6:
                    x.fill = fill(gfill[l["grade"]]); x.font = Font(bold=True)
                if status == "н" and i == 5:
                    x.fill = fill("FFE4E4")
                if status == "уваж." and i == 5:
                    x.fill = fill("FFF3CD")
            ri += 1
    for i, w in enumerate([24, 12, 32, 8, 10, 10, 28, 12], 1):
        wsA.column_dimensions[L(i)].width = w
    wsA.freeze_panes = "A4"
    if ri > 4:
        wsA.auto_filter.ref = f"A3:H{ri-1}"

    # ===== Sheet 3: Итоги =====
    wsF = wb.create_sheet("Итоговые аттестации")
    wsF["A1"] = f'Итоговые: {st["name"]}'
    wsF["A1"].font = title_font
    fh = ["Предмет", "Аттестация", "Период", "Оценка", "Заметка", "Обновлено"]
    for i, h in enumerate(fh, 1):
        x = wsF.cell(3, i, h); x.font = head_font; x.fill = fill("14213D"); x.border = box; x.alignment = center
    ri = 4
    for sub in subjects:
        kinds = get_final_kinds(d, st["group_id"], sub["id"])
        frows = {r["kind"]: r for r in d.execute(
            "SELECT * FROM final_grades WHERE student_id=? AND subject_id=?", (uid, sub["id"]))}
        for k, lab in kinds.items():
            g = frows.get(k)
            per = period_for(d, st["group_id"], sub["id"], k)
            vals = [sub["name"], lab, per.get("label") or "",
                    (g["grade"] if g else "") or "—",
                    (g["note"] if g else "") or "",
                    (g["updated"] if g else "") or ""]
            for i, v in enumerate(vals, 1):
                x = wsF.cell(ri, i, v); x.border = box
                if g and g["grade"] and i == 4:
                    x.fill = fill("B8F0D3"); x.font = Font(bold=True)
            ri += 1
    for i, w in enumerate([24, 28, 22, 12, 30, 16], 1):
        wsF.column_dimensions[L(i)].width = w

    # ===== Sheet 4: ДЗ =====
    wsH = wb.create_sheet("Задания")
    wsH["A1"] = f'Задания: {st["name"]}'
    wsH["A1"].font = title_font
    hh = ["Предмет", "Задание", "Срок", "Статус", "Выполнено"]
    for i, h in enumerate(hh, 1):
        x = wsH.cell(3, i, h); x.font = head_font; x.fill = fill("14213D"); x.border = box; x.alignment = center
    ri = 4
    if st["group_id"]:
        hws = d.execute("""SELECT h.title, h.due, h.created, s.name sname,
            (SELECT 1 FROM hw_done x WHERE x.hw_id=h.id AND x.student_id=?) done
            FROM homework h JOIN subjects s ON s.id=h.subject_id
            WHERE h.group_id=? ORDER BY h.due DESC""", (uid, st["group_id"])).fetchall()
        for h in hws:
            vals = [h["sname"], h["title"], h["due"] or "", "сдано" if h["done"] else "не сдано", "✓" if h["done"] else ""]
            for i, v in enumerate(vals, 1):
                x = wsH.cell(ri, i, v); x.border = box
                if not h["done"] and i == 4:
                    x.fill = fill("FFF3CD")
            ri += 1
    for i, w in enumerate([22, 40, 14, 12, 12], 1):
        wsH.column_dimensions[L(i)].width = w

    # ===== per-subject sheets (capped) =====
    for sub in subjects[:20]:
        safe_title = re.sub(r'[\\/*?:\[\]]', "_", sub["name"])[:28] or "Предмет"
        # unique sheet name
        base, n = safe_title, 1
        while safe_title in [s.title for s in wb.worksheets]:
            n += 1
            safe_title = f"{base[:25]}_{n}"
        ws2 = wb.create_sheet(title=safe_title)
        ws2["A1"] = f'{st["name"]} — {sub["name"]}'
        ws2["A1"].font = Font(size=12, bold=True)
        tname = sub["tname"] if "tname" in sub.keys() else ""
        ws2["A2"] = f'Преподаватели: {tname or "—"}'
        for i, h in enumerate(["Дата", "Тема", "Статус", "Оценка", "Причина"], 1):
            x = ws2.cell(4, i, h); x.font = head_font; x.fill = fill("14213D"); x.border = box
        lessons = d.execute("""SELECT l.id, l.day, l.topic, m.present, m.grade, m.excuse FROM lessons l
            LEFT JOIN marks m ON m.lesson_id=l.id AND m.student_id=?
            WHERE l.group_id=? AND l.subject_id=? ORDER BY l.day""", (uid, st["group_id"], sub["id"])).fetchall()
        for ri, l in enumerate(lessons, 5):
            ws2.cell(ri, 1, f'{l["day"][8:10]}.{l["day"][5:7]}' if l["day"] else "").border = box
            ws2.cell(ri, 2, l["topic"] or "").border = box
            status = "—"
            if l["grade"]: status = "оценка"
            elif l["present"] == 2: status = "уваж."
            elif l["present"] == 0: status = "н"
            elif l["present"] == 1: status = "был"
            ws2.cell(ri, 3, status).border = box
            gx = ws2.cell(ri, 4, l["grade"] or ""); gx.border = box
            if l["grade"] and l["grade"] in gfill:
                gx.fill = fill(gfill[l["grade"]]); gx.font = Font(bold=True)
            ws2.cell(ri, 5, (l["excuse"] or "") if l["present"] == 2 else "").border = box
        for i in range(1, 6):
            ws2.column_dimensions[L(i)].width = [12, 30, 12, 10, 36][i-1]
        stt = stats(lessons)
        fr = 5 + len(lessons) + 1
        ws2.cell(fr, 1, f'Средний: {stt["avg"] or "—"}; пропуски: {stt["missed"]}; посещ.: {stt["pct"] if stt["pct"] is not None else "—"}%')
        kinds = get_final_kinds(d, st["group_id"], sub["id"])
        frows = {r["kind"]: r for r in d.execute(
            "SELECT * FROM final_grades WHERE student_id=? AND subject_id=?", (uid, sub["id"]))}
        ws2.cell(fr + 2, 1, "Итоговые аттестации").font = Font(bold=True)
        for i, (k, lab) in enumerate(kinds.items()):
            per = period_for(d, st["group_id"], sub["id"], k)
            g = frows.get(k)
            ws2.cell(fr + 3 + i, 1, lab).border = box
            ws2.cell(fr + 3 + i, 2, per.get("label") or "").border = box
            ws2.cell(fr + 3 + i, 3, (g["grade"] if g else "") or "—").border = box
            ws2.cell(fr + 3 + i, 4, (g["note"] if g else "") or "").border = box

    buf = io.BytesIO(); wb.save(buf)
    safe = "".join(c if c.isalnum() or c in "-_" else "_" for c in st["name"])[:40]
    return Response(buf.getvalue(), mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename=student_{safe}.xlsx"})

@bp.route("/export/all")
@role("admin")
def export_all():
    """Выгрузка сводки по всем студентам."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter as L
    d = db()
    wb = Workbook()
    fill = lambda h: PatternFill("solid", fgColor=h)
    thin = Side(style="thin", color="D9DEEC"); box = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")
    ws = wb.active; ws.title = "Все студенты"
    ws["A1"] = f'Сводка по всем студентам. Выгружено {datetime.now():%d.%m.%Y %H:%M}'
    ws["A1"].font = Font(size=14, bold=True, color="14213D")
    heads = ["Студент", "Группа", "Предмет", "Ср. балл", "Оценок", "Был", "Не был", "Уваж.", "Пропуски", "Посещаемость %", "Доп. баллы", "Итоги"]
    for i, h in enumerate(heads, 1):
        x = ws.cell(3, i, h); x.font = Font(bold=True, color="FFFFFF"); x.fill = fill("14213D"); x.border = box; x.alignment = center
    ri = 4
    students = d.execute("""SELECT u.id, u.name, u.group_id, g.name gname FROM users u
        LEFT JOIN groups g ON g.id=u.group_id WHERE u.role='student' ORDER BY g.name, u.name""").fetchall()
    for st in students:
        if not st["group_id"]:
            ws.cell(ri, 1, st["name"]).border = box
            ws.cell(ri, 2, "—").border = box
            ri += 1
            continue
        subjects = d.execute("""SELECT s.id, s.name FROM subjects s
            WHERE s.id IN (SELECT subject_id FROM assignments WHERE group_id=?) ORDER BY s.name""", (st["group_id"],)).fetchall()
        for sub in subjects:
            cells = d.execute("""SELECT m.present, m.grade, m.excuse FROM lessons l
                LEFT JOIN marks m ON m.lesson_id=l.id AND m.student_id=?
                WHERE l.group_id=? AND l.subject_id=?""", (st["id"], st["group_id"], sub["id"])).fetchall()
            stt = stats(cells)
            bp = d.execute("SELECT balance FROM bonus_points WHERE student_id=? AND subject_id=?", (st["id"], sub["id"])).fetchone()
            grades_n = sum(1 for c in cells if c and c["grade"])
            was = sum(1 for c in cells if c and c["present"] == 1)
            kinds = get_final_kinds(d, st["group_id"], sub["id"])
            frows = d.execute("SELECT kind, grade FROM final_grades WHERE student_id=? AND subject_id=?", (st["id"], sub["id"])).fetchall()
            by_k = {r["kind"]: r["grade"] for r in frows if r["grade"]}
            finals_txt = "; ".join(f"{kinds.get(k,k)}:{g}" for k, g in by_k.items()) or ""
            vals = [st["name"], st["gname"] or "—", sub["name"], stt["avg"], grades_n, was,
                    stt.get("unexcused") or 0, stt.get("excused") or 0, stt["missed"],
                    stt["pct"], bp["balance"] if bp else 0, finals_txt]
            for i, v in enumerate(vals, 1):
                x = ws.cell(ri, i, v); x.border = box; x.alignment = center if i > 2 else Alignment(horizontal="left")
            ri += 1
    for i, w in enumerate([22, 12, 24, 11, 9, 8, 9, 9, 10, 12, 11, 36], 1):
        ws.column_dimensions[L(i)].width = w
    buf = io.BytesIO(); wb.save(buf)
    return Response(buf.getvalue(), mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": "attachment; filename=all_students.xlsx"})


# ---------- Админ tools ----------
# ---------- Админ: служебные действия ----------
@bp.post("/admin/tools")
@role("admin")
def admin_tools():
    d, f = db(), request.form
    act = f.get("act")
    try:
        if act == "clear_marks_group":
            gid = int(f["group_id"])
            d.execute("""DELETE FROM marks WHERE student_id IN
              (SELECT id FROM users WHERE role='student' AND group_id=?)""", (gid,))
            d.execute("""DELETE FROM lessons WHERE group_id=?""", (gid,))
            flash("Оценки и занятия группы очищены")
        elif act == "clear_marks_student":
            sid = int(f["student_id"])
            d.execute("DELETE FROM marks WHERE student_id=?", (sid,))
            d.execute("DELETE FROM bonus_points WHERE student_id=?", (sid,))
            d.execute("DELETE FROM final_grades WHERE student_id=?", (sid,))
            flash("Данные студента (оценки, доп. баллы, итоги) очищены")
        elif act == "clear_all_notifications":
            d.execute("DELETE FROM notifications")
            flash("Все уведомления удалены")
        elif act == "cleanup_bells":
            n = 0
            for b in d.execute("SELECT id,num,start,end FROM bells").fetchall():
                if b["num"] > 12 or b["num"] < 1:
                    used = d.execute("SELECT COUNT(*) c FROM schedule WHERE bell_id=?", (b["id"],)).fetchone()["c"]
                    if used == 0:
                        d.execute("DELETE FROM bells WHERE id=?", (b["id"],)); n += 1
                elif b["start"]=="00:00" and b["end"]=="00:00":
                    defaults = {1:("09:00","10:35"),2:("10:55","12:30"),3:("13:10","14:45"),
                                4:("14:55","16:30"),5:("16:40","18:15"),6:("18:25","20:00")}
                    if 1 <= b["num"] <= 6 and b["num"] in defaults:
                        s,e = defaults[b["num"]]
                        d.execute("UPDATE bells SET start=?,end=? WHERE id=?", (s,e,b["id"])); n += 1
                    else:
                        used = d.execute("SELECT COUNT(*) c FROM schedule WHERE bell_id=?", (b["id"],)).fetchone()["c"]
                        if used == 0:
                            d.execute("DELETE FROM bells WHERE id=?", (b["id"],)); n += 1
            flash(f"Звонки: исправлено записей — {n}")
        elif act == "logout_hint":
            # Flask sessions are cookie-based; force password rotate is safest "logout all"
            flash("Сессии cookie-based: смените SECRET_KEY или пароли пользователей, чтобы инвалидировать входы")
        elif act == "personal_announce":
            uid = int(f["user_id"])
            title = f["title"].strip()
            body = f["body"].strip()
            if title and body:
                try:
                    d.execute("INSERT INTO announcements(author_id,group_id,title,body,created,target_user_id) VALUES(?,?,?,?,?,?)",
                              (session["uid"], None, title, body, datetime.now().strftime("%d.%m.%Y %H:%M"), uid))
                except Exception:
                    d.execute("INSERT INTO announcements(author_id,group_id,title,body,created) VALUES(?,?,?,?,?)",
                              (session["uid"], None, title, body, datetime.now().strftime("%d.%m.%Y %H:%M")))
                notify(uid, "Личное: " + title, body[:200], url_for("notifications_page"))
                flash("Личное объявление отправлено")
            else:
                flash("Укажите заголовок и текст")
        else:
            flash("Неизвестное действие")
        d.commit()
    except Exception as e:
        d.rollback()
        flash(f"Ошибка: {e}")
    return redirect(url_for("admin"))



# ---------- Backup ----------
@bp.route("/admin/backup")
@role("admin")
def admin_backup():
    import zipfile, csv
    d = db()
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        try:
            d.execute("PRAGMA wal_checkpoint(FULL)")
        except Exception: pass
        z.write(DB, "campus.db")
        out = io.StringIO(); w = csv.writer(out)
        w.writerow(["id","login","name","role","group","email","phone","birth_date","student_id_number",
                    "parent_name","parent_phone","address","notes","active"])
        for u in d.execute("""SELECT u.*, g.name gname FROM users u LEFT JOIN groups g ON g.id=u.group_id
                              ORDER BY u.role, g.name, u.name"""):
            keys = u.keys()
            w.writerow([u["id"], u["login"], u["name"], u["role"], u["gname"] or "",
                        u["email"] if "email" in keys else "", u["phone"] if "phone" in keys else "",
                        u["birth_date"] if "birth_date" in keys else "",
                        u["student_id_number"] if "student_id_number" in keys else "",
                        u["parent_name"] if "parent_name" in keys else "",
                        u["parent_phone"] if "parent_phone" in keys else "",
                        u["address"] if "address" in keys else "",
                        u["notes"] if "notes" in keys else "",
                        u["active"] if "active" in keys else 1])
        z.writestr("users.csv", out.getvalue())
        out2 = io.StringIO(); w2 = csv.writer(out2)
        w2.writerow(["student","group","subject","day","present","grade","excuse"])
        for m in d.execute("""SELECT u.name uname, g.name gname, s.name sname, l.day, m.present, m.grade,
                              COALESCE(m.excuse,'') excuse
                              FROM marks m
                              JOIN users u ON u.id=m.student_id
                              JOIN lessons l ON l.id=m.lesson_id
                              JOIN groups g ON g.id=l.group_id
                              JOIN subjects s ON s.id=l.subject_id
                              ORDER BY g.name, s.name, l.day, u.name"""):
            w2.writerow([m["uname"], m["gname"], m["sname"], m["day"], m["present"], m["grade"], m["excuse"]])
        z.writestr("marks.csv", out2.getvalue())
    buf.seek(0)
    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    return Response(buf.getvalue(), mimetype="application/zip",
                    headers={"Content-Disposition": f"attachment; filename=campus_backup_{stamp}.zip"})

